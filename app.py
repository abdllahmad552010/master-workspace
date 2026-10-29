#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Abdullah Master Workspace
-------------------------
ثلاث مشاريع في تطبيق واحد: منظم أحداث (أسبوع / شهر) + مولّد سلايدات من الكتب + متابعة المنهج.
التنقل بينهم من قايمة ⋮ فوق.

التشغيل:
    pip install flask pypdf python-docx python-pptx
    python3 app.py            ثم افتح  http://localhost:5000

اختياري:
    ANTHROPIC_API_KEY=...   لتفعيل تحسين التلخيص بواسطة Claude
    HOST=0.0.0.0            لو عايز تفتحه من موبايل على نفس الشبكة
"""
import importlib.util
import io
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import traceback
import unicodedata
import urllib.request
import uuid
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

from flask import Flask, Response, jsonify, request, send_file

# --------------------------------------------------------------------------
# الإعدادات
# --------------------------------------------------------------------------
BASE_DIR = os.path.expanduser("~/DragonWorkspace")
DATA_FILE = os.path.join(BASE_DIR, "app_data.json")
UPLOAD_DIR = os.path.join(BASE_DIR, ".uploads_tmp")
os.makedirs(BASE_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

MAX_FILE_BYTES = 2 * 1024 ** 3          # 2GB
MAX_TEXT_FILE_READ = 150 * 1024 ** 2    # أقصى قراءة لملف نصي
ALLOWED_EXT = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024   # حجم الطلب الواحد (الشريحة)


def have(mod):
    try:
        return importlib.util.find_spec(mod) is not None
    except Exception:
        return False


# --------------------------------------------------------------------------
# تخزين البيانات (الأحداث + آخر عرض)
# --------------------------------------------------------------------------
DATA_LOCK = threading.RLock()
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _fresh():
    return {"events": {}, "notes": "", "deck": None, "study": {"progress": {}, "custom": {}}}


def load_data():
    with DATA_LOCK:
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                d = json.load(f)
        except (OSError, ValueError):
            d = {}
    if not isinstance(d, dict):
        d = {}
    base = _fresh()
    base.update(d)
    if not isinstance(base["events"], dict):
        base["events"] = {}
    base["study"] = clean_study(base.get("study"))
    return base


def save_data(d):
    with DATA_LOCK:
        tmp = DATA_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=2)
        os.replace(tmp, DATA_FILE)


def clean_events(ev):
    out = {}
    if not isinstance(ev, dict):
        return out
    for date, items in ev.items():
        if not isinstance(date, str) or not DATE_RE.match(date) or not isinstance(items, list):
            continue
        rows = []
        for it in items[:200]:
            if not isinstance(it, dict):
                continue
            title = str(it.get("title", "")).strip()[:300]
            if not title:
                continue
            time_s = str(it.get("time", "08:00"))[:5]
            if not re.match(r"^\d{2}:\d{2}$", time_s):
                time_s = "08:00"
            try:
                remind = max(0, min(1440, int(it.get("remind", 0))))
            except (TypeError, ValueError):
                remind = 0
            rows.append({
                "id": str(it.get("id") or uuid.uuid4().hex[:10])[:24],
                "title": title,
                "time": time_s,
                "remind": remind,
                "done": bool(it.get("done", False)),
            })
        if rows:
            out[date] = rows
    return out


def clean_study(st):
    out = {"progress": {}, "custom": {}}
    if not isinstance(st, dict):
        return out
    prog = st.get("progress")
    if isinstance(prog, dict):
        for k, v in list(prog.items())[:5000]:
            if isinstance(k, str) and len(k) <= 80 and v:
                out["progress"][k] = True
    cust = st.get("custom")
    if isinstance(cust, dict):
        for subj, items in list(cust.items())[:50]:
            if not isinstance(subj, str) or len(subj) > 40 or not isinstance(items, list):
                continue
            rows = []
            for it in items[:500]:
                if isinstance(it, dict):
                    t = str(it.get("title", "")).strip()[:300]
                    i = str(it.get("id", ""))[:60]
                    if t and i:
                        rows.append({"id": i, "title": t})
            if rows:
                out["custom"][subj] = rows
    return out


def clean_deck(deck):
    if not isinstance(deck, dict):
        return None
    slides = []
    for s in (deck.get("slides") or [])[:500]:
        if not isinstance(s, dict):
            continue
        slides.append({
            "type": "title" if s.get("type") == "title" else "content",
            "title": str(s.get("title", ""))[:300],
            "bullets": [str(b)[:600] for b in (s.get("bullets") or [])[:40]],
        })
    if not slides:
        return None
    return {"title": str(deck.get("title", "عرض"))[:200], "slides": slides}


# --------------------------------------------------------------------------
# معالجة اللغة: تقسيم الجمل، الكلمات، الكلمات الشائعة
# --------------------------------------------------------------------------
AR_DIAC = re.compile(r"[ؗ-ًؚ-ْـ]")
WORD_RE = re.compile(r"[\w؀-ۿ]+", re.UNICODE)
AR_CHARS = re.compile(r"[؀-ۿ]")
SENT_SPLIT = re.compile(r"(?<=[\.\!\?؟…۔])\s+|\n+")
BULLET_START = re.compile(r"^(?:[-•*▪●–]\s+|\d{1,3}[\.\)]\s+|[A-Za-z][\.\)]\s+)")
HEAD_KEY = re.compile(
    r"^(chapter|section|unit|lecture|part|module|الفصل|الباب|المحاضرة|الوحدة|القسم|الجزء|الدرس|الموضوع)\b",
    re.I)
NUM_HEAD = re.compile(r"^\d{1,2}(\.\d{1,2}){0,3}[\.\)]?\s+\S")


def norm_ar(w):
    w = AR_DIAC.sub("", w.lower())
    w = re.sub("[أإآ]", "ا", w)
    return w.replace("ى", "ي").replace("ة", "ه")


_STOP_AR = """في من على الى إلى عن مع هذا هذه ذلك تلك التي الذي الذين التى هو هي هم هما نحن انا أنا انت أنت كان كانت
يكون تكون ان أن إن إذا اذا ما لا لم لن قد كل بعض غير بين حتى عند عندما حيث او أو ثم كما لكن لكنه أي اي هناك هنا بعد قبل
خلال فقط جدا ايضا أيضا وقد وكان وهو وهي وفي ومن وعلى والى وان وأن ولا ولم ولن به بها بهم له لها لهم منه منها منهم فيه
فيها فيهم عليه عليها عليهم اليه اليها كذلك يمكن بما لما مما اما أما اذ إذ هل كي لو ليس ليست يتم تم ذات احد""".split()
_STOP_EN = """the a an and or of to in on at by for with from as is are was were be been being it its this that these those
which who whom what when where why how not no nor but if then than so such can could may might will would shall should do
does did done has have had having he she they we you i his her their our your them him also into over under about between
through during before after above below up down out off again further once only own same too very just more most other some
any each few both all there here""".split()
STOP = {norm_ar(w) for w in _STOP_AR} | set(_STOP_EN)

LABELS = {
    True: {"agenda": "المحاور الرئيسية", "end": "الخلاصة والنقاط الأهم", "auto": "ملخص تلقائي",
           "pages": "صفحة", "words": "كلمة", "cont": "تابع"},
    False: {"agenda": "Agenda", "end": "Key takeaways", "auto": "Auto summary",
            "pages": "pages", "words": "words", "cont": "cont."},
}


def arabic_ratio(text):
    sample = text[:200000]
    letters = sum(1 for c in sample if c.isalpha())
    if not letters:
        return 0.0
    return len(AR_CHARS.findall(sample)) / letters


def shorten(s, n=170):
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"^\s*(?:\d{1,3}[\.\)]|[-•*▪●–])\s*", "", s)
    if len(s) <= n:
        return s
    cut = s[:n].rsplit(" ", 1)[0]
    return (cut or s[:n]).rstrip(" ,;:،؛-") + "…"


class Analyzer:
    """تقطيع الكلمات مع جذر بسيط للعربي والإنجليزي + حفظ الشكل الأصلي للكلمة."""

    def __init__(self):
        self.surf = defaultdict(Counter)

    @staticmethod
    def stem(w):
        if AR_CHARS.search(w):
            w = norm_ar(w)
            for p in ("وال", "فال", "بال", "كال", "لل", "ال"):
                if w.startswith(p) and len(w) - len(p) >= 3:
                    w = w[len(p):]
                    break
            return w
        if len(w) > 4 and w.endswith("s") and not w.endswith("ss"):
            w = w[:-1]
        return w

    def tokenize(self, s):
        out = []
        for w in WORD_RE.findall(AR_DIAC.sub("", s)):
            lw = w.lower()
            if len(lw) < 3 or lw.isdigit() or "_" in lw:
                continue
            st = self.stem(lw)
            if len(st) < 3 or st in STOP:
                continue
            self.surf[st][AR_DIAC.sub("", w)] += 1
            out.append(st)
        return out

    def surface(self, st):
        c = self.surf.get(st)
        return c.most_common(1)[0][0] if c else st


def is_heading(line):
    l = line.strip()
    if l.startswith("# "):
        return l[2:].strip() or None
    if len(l) < 3 or len(l) > 80 or l[-1] in ".،؛,;:!?؟" or re.search(r"\.{4,}", l):
        return None
    if HEAD_KEY.match(l) or NUM_HEAD.match(l):
        return l
    return None


def split_sentences(par):
    out = []
    for s in SENT_SPLIT.split(par):
        s = s.strip(" \t-•–—*▪●")
        if len(s) < 25:
            continue
        if re.search(r"https?://|www\.", s):
            continue
        junk = sum(1 for c in s if not (c.isalpha() or c.isspace()))
        if junk / len(s) > 0.35:
            continue
        if len(s) > 420:
            buf = ""
            for p in re.split(r"(?<=[،؛;:,])\s+", s):
                if buf and len(buf) + len(p) > 260:
                    out.append(buf.strip())
                    buf = p
                else:
                    buf = (buf + " " + p).strip()
            if buf:
                out.append(buf)
        else:
            out.append(s)
    return out


def parse_sections(text):
    """يحوّل النص إلى أقسام: {'title', 'sents', 'chars'} باستخدام العناوين لو موجودة."""
    sections, cur, buf = [], {"title": None, "paras": []}, []

    def flush():
        if buf:
            cur["paras"].append(" ".join(buf))
            buf.clear()

    def close():
        sents = []
        for p in cur["paras"]:
            sents.extend(split_sentences(p))
        if sents:
            sections.append({"title": cur["title"], "sents": sents,
                             "chars": sum(len(x) for x in sents)})

    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            flush()
            continue
        h = is_heading(line)
        if h:
            flush()
            if cur["paras"]:
                close()
            cur = {"title": h, "paras": []}
            continue
        if BULLET_START.match(line):
            flush()
        if buf and buf[-1].endswith("-") and len(buf[-1]) > 2 and buf[-1][-2].isalpha():
            buf[-1] = buf[-1][:-1] + line
        else:
            buf.append(line)
    flush()
    close()

    merged = []
    for s in sections:
        if merged and s["chars"] < 200:
            merged[-1]["sents"].extend(s["sents"])
            merged[-1]["chars"] += s["chars"]
        else:
            merged.append(s)
    return merged


def _merge(a, b):
    a["sents"] = a["sents"] + b["sents"]
    a["chars"] += b["chars"]
    if not a["title"]:
        a["title"] = b["title"]


def fit_sections(sections, budget, cont_label):
    sections = [dict(s) for s in sections]
    if len(sections) > budget:
        total = sum(s["chars"] for s in sections)
        tgt = total / budget
        groups, cur = [], None
        for s in sections:
            if cur is None:
                cur = s
            elif cur["chars"] < tgt * 0.85:
                _merge(cur, s)
            else:
                groups.append(cur)
                cur = s
        groups.append(cur)
        while len(groups) > budget:
            i = min(range(len(groups) - 1), key=lambda k: groups[k]["chars"] + groups[k + 1]["chars"])
            _merge(groups[i], groups[i + 1])
            del groups[i + 1]
        sections = groups
    while len(sections) < budget:
        i = max(range(len(sections)), key=lambda k: sections[k]["chars"])
        sec = sections[i]
        if sec["chars"] < 1500 or len(sec["sents"]) < 6:
            break
        half, acc, cut = sec["chars"] / 2, 0, 1
        for j, st in enumerate(sec["sents"]):
            acc += len(st)
            if acc >= half:
                cut = max(1, min(j + 1, len(sec["sents"]) - 1))
                break
        a = {"title": sec["title"], "sents": sec["sents"][:cut]}
        b = {"title": (sec["title"] + " (" + cont_label + ")") if sec["title"] else None,
             "sents": sec["sents"][cut:]}
        for x in (a, b):
            x["chars"] = sum(len(t) for t in x["sents"])
        sections[i:i + 1] = [a, b]
    return sections


def _jaccard(a, b):
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def pick(sents, toks, tf, idf, k, avoid=()):
    scores = []
    for i, t in enumerate(toks):
        u = set(t)
        if len(u) < 3:
            scores.append(0.0)
            continue
        sc = sum(tf[w] * idf.get(w, 1.0) for w in u) / math.sqrt(len(u) + 4)
        if i == 0:
            sc *= 1.2
        scores.append(sc)
    chosen, chosen_sets = [], []
    for i in sorted(range(len(sents)), key=lambda x: scores[x], reverse=True):
        if len(chosen) >= k or scores[i] <= 0:
            break
        u = set(toks[i])
        if any(_jaccard(u, c) > 0.55 for c in chosen_sets) or any(_jaccard(u, c) > 0.5 for c in avoid):
            continue
        chosen.append(i)
        chosen_sets.append(u)
    chosen.sort()
    return chosen


# --------------------------------------------------------------------------
# Claude (اختياري)
# --------------------------------------------------------------------------
def ai_available():
    return bool(os.environ.get("ANTHROPIC_API_KEY"))


def ai_summarize(title, text, ar):
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        return None
    lang = "العربية" if ar else "English"
    prompt = (
        "لخّص القسم التالي من كتاب/مذكرة لسلايد عرض تقديمي.\n"
        f"اللغة المطلوبة: {lang} (نفس لغة النص).\n"
        'رُد بـ JSON فقط بهذا الشكل: {"title": "عنوان قصير", "bullets": ["نقطة", "نقطة", "نقطة"]}\n'
        "من 3 إلى 5 نقاط، كل نقطة أقل من 150 حرفاً، بدون أي كلام خارج الـ JSON.\n\n"
        f"عنوان القسم: {title or '(غير محدد)'}\n\nالنص:\n{text}"
    )
    body = json.dumps({
        "model": os.environ.get("ANTHROPIC_MODEL", "claude-sonnet-5-5"),
        "max_tokens": 800,
        "messages": [{"role": "user", "content": prompt}],
    }).encode("utf-8")
    req = urllib.request.Request(
        "https://api.anthropic.com/v1/messages", data=body,
        headers={"content-type": "application/json", "x-api-key": key,
                 "anthropic-version": "2023-06-01"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            data = json.loads(r.read().decode("utf-8"))
        txt = "".join(b.get("text", "") for b in data.get("content", []) if b.get("type") == "text")
        obj = json.loads(re.search(r"\{.*\}", txt, re.S).group(0))
        bullets = [str(b).strip() for b in obj.get("bullets", []) if str(b).strip()]
        if not bullets:
            return None
        return {"title": str(obj.get("title", "")).strip(), "bullets": bullets[:6]}
    except Exception:
        return None


# --------------------------------------------------------------------------
# بناء العرض (Deck)
# --------------------------------------------------------------------------
def build_deck(text, title, target, use_ai, cb, meta=None):
    meta = meta or {}
    if re.search(r"[ﭐ-﷿ﹰ-﻿]", text):
        text = unicodedata.normalize("NFKC", text)
    ar = arabic_ratio(text) > 0.3
    L = LABELS[ar]

    sections = parse_sections(text)
    if not sections:
        raise ValueError("ما قدرتش ألاقي جمل كافية في الملف لتلخيصها.")
    total_chars = sum(s["chars"] for s in sections)
    if not target:
        target = min(40, max(8, 6 + total_chars // 4000))
    target = max(6, min(int(target), 80))
    budget = max(1, min(target - 3, total_chars // 400 or 1))
    sections = fit_sections(sections, budget, L["cont"])

    an = Analyzer()
    cb(0.05, "تحليل الكلمات المهمة...")
    for s in sections:
        s["tok"] = [an.tokenize(x) for x in s["sents"]]
        s["tf"] = Counter(w for t in s["tok"] for w in t)
    n = len(sections)
    df = Counter()
    for s in sections:
        df.update(s["tf"].keys())
    idf = {w: math.log(1 + n / c) for w, c in df.items()}

    use_ai = bool(use_ai and ai_available())
    slides, used_sets, used_idx = [], [], []
    for si, s in enumerate(sections):
        cb(0.1 + 0.8 * si / n, f"تلخيص القسم {si + 1} من {n}" + (" بواسطة Claude" if use_ai else ""))
        k = 3 if s["chars"] < 1500 else 4 if s["chars"] < 6000 else 5
        idxs = pick(s["sents"], s["tok"], s["tf"], idf, k)
        if not idxs:
            idxs = list(range(min(k, len(s["sents"]))))
        bullets = [shorten(s["sents"][i]) for i in idxs]
        for i in idxs:
            used_sets.append(set(s["tok"][i]))
            used_idx.append((si, i))
        if s["title"]:
            sl_title = shorten(s["title"].lstrip("# ").strip(" :-–"), 90)
        else:
            top = sorted(s["tf"], key=lambda w: s["tf"][w] * idf.get(w, 1), reverse=True)[:3]
            sl_title = " • ".join(an.surface(w) for w in top) or f"{si + 1}"
        if use_ai:
            res = ai_summarize(s["title"], " ".join(s["sents"])[:9000], ar)
            if res:
                bullets = [shorten(b, 190) for b in res["bullets"]]
                sl_title = shorten(res["title"], 90) or sl_title
        slides.append({"type": "content", "title": sl_title, "bullets": bullets})

    pages = meta.get("pages")
    words = total_chars // 6
    sub = f"{L['auto']} · {pages} {L['pages']}" if pages else f"{L['auto']} · ~{words} {L['words']}"
    deck = [{"type": "title", "title": title, "bullets": [sub]}]
    if len(slides) >= 3:
        deck.append({"type": "content", "title": L["agenda"],
                     "bullets": [shorten(x["title"], 80) for x in slides[:12]]})
    deck.extend(slides)

    # الخلاصة: أهم جمل في الملف كله (غير المستخدمة)
    cb(0.93, "كتابة الخلاصة...")
    gtf = Counter()
    for s in sections:
        gtf.update(s["tf"])
    cand = []
    for si, s in enumerate(sections):
        for i, t in enumerate(s["tok"]):
            u = set(t)
            if len(u) < 4 or (si, i) in used_idx:
                continue
            sc = sum(gtf[w] * idf.get(w, 1.0) for w in u) / math.sqrt(len(u) + 4)
            cand.append((sc, si, i))
    cand.sort(reverse=True)
    end_bullets, end_sets = [], []
    for sc, si, i in cand:
        u = set(sections[si]["tok"][i])
        if any(_jaccard(u, c) > 0.5 for c in used_sets + end_sets):
            continue
        end_bullets.append(shorten(sections[si]["sents"][i]))
        end_sets.append(u)
        if len(end_bullets) >= 5:
            break
    if end_bullets and len(slides) >= 2:
        deck.append({"type": "content", "title": L["end"], "bullets": end_bullets})
    cb(1.0, "تم")
    return {"title": title, "slides": deck}


# --------------------------------------------------------------------------
# استخراج النص من الملفات
# --------------------------------------------------------------------------
def _decode(raw):
    for enc in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("latin-1", errors="replace")


def clean_pages(pages):
    """يشيل رؤوس وأذيال الصفحات المتكررة وأرقام الصفحات."""
    n = len(pages)
    split = [[l.strip() for l in p.split("\n")] for p in pages]
    repeated = set()
    if n >= 4:
        cnt = Counter()
        for lines in split:
            cnt.update({l for l in lines if 0 < len(l) < 80})
        repeated = {l for l, c in cnt.items() if c >= max(3, n * 0.3)}
    out = []
    for lines in split:
        keep = [l for l in lines if l not in repeated and not re.fullmatch(r"[\-–—\s]*\d{1,4}[\-–—\s]*", l)]
        out.append("\n".join(keep))
    return out


_OCR_LANGS = None


def ocr_available():
    return bool(shutil.which("tesseract") and shutil.which("pdftoppm"))


def ocr_langs():
    """يرجع (لغة_tesseract, هل_العربي_مثبت)."""
    global _OCR_LANGS
    if _OCR_LANGS is None:
        try:
            out = subprocess.run([shutil.which("tesseract"), "--list-langs"], capture_output=True,
                                 text=True, timeout=20).stdout.split()
        except Exception:
            out = []
        langs = [x for x in ("ara", "eng") if x in out] or ["eng"]
        _OCR_LANGS = ("+".join(langs), "ara" in langs)
    return _OCR_LANGS


def ocr_page(path, pageno, lang):
    """يحوّل صفحة PDF لصورة ثم يقرأها بـ tesseract."""
    with tempfile.TemporaryDirectory(dir=UPLOAD_DIR) as td:
        prefix = os.path.join(td, "p")
        subprocess.run([shutil.which("pdftoppm"), "-f", str(pageno), "-l", str(pageno), "-r", "200",
                        "-gray", "-png", path, prefix], check=True, timeout=180,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        imgs = [f for f in os.listdir(td) if f.endswith(".png")]
        if not imgs:
            return ""
        env = dict(os.environ, OMP_THREAD_LIMIT="1")
        r = subprocess.run([shutil.which("tesseract"), os.path.join(td, imgs[0]), "stdout", "-l", lang,
                            "--psm", "3"], capture_output=True, timeout=300, env=env)
        return r.stdout.decode("utf-8", "replace")


def extract_pdf(path, cb, ocr=True):
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise RuntimeError("مكتبة قراءة PDF غير مثبتة. نفّذ: pip install pypdf")
    r = PdfReader(path)
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            raise RuntimeError("الملف محمي بكلمة سر.")
    n = len(r.pages)
    pages = []
    for i, p in enumerate(r.pages):
        try:
            pages.append(p.extract_text() or "")
        except Exception:
            pages.append("")
        if i % 3 == 0 or i == n - 1:
            cb((i + 1) / n * 150, 1000, f"قراءة الصفحة {i + 1} من {n}")

    meta = {"pages": n}
    scanned = [i for i, t in enumerate(pages) if len(t.strip()) < 30]
    meta["scanned"] = len(scanned)
    if scanned and ocr and ocr_available():
        lang, has_ara = ocr_langs()
        meta["ocr"] = True
        if not has_ara:
            meta["warn"] = "الـ OCR اشتغل بالإنجليزي فقط. للعربي ثبّت: sudo apt install tesseract-ocr-ara"
        workers = max(1, min(4, os.cpu_count() or 1))
        total = len(scanned)
        done = 0
        with ThreadPoolExecutor(max_workers=workers) as ex:
            futs = {ex.submit(ocr_page, path, i + 1, lang): i for i in scanned}
            for fu in as_completed(futs):
                i = futs[fu]
                try:
                    pages[i] = fu.result()
                except Exception:
                    pages[i] = ""
                done += 1
                cb(150 + done / total * 850, 1000, f"OCR: قراءة صفحة ممسوحة {done} من {total} (قد يستغرق وقتاً)")
    else:
        cb(1000, 1000, "تمت قراءة الملف")
    return "\n".join(clean_pages(pages)), meta


def extract_docx(path, cb):
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise RuntimeError("مكتبة Word غير مثبتة. نفّذ: pip install python-docx")
    d = docx.Document(path)
    body = list(d.element.body.iterchildren())
    lines = []
    for i, child in enumerate(body):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                lines.append("")
                continue
            sname = (p.style.name if p.style is not None else "") or ""
            if sname.lower().startswith(("heading", "title")) or sname.startswith("عنوان"):
                lines.append("# " + t)
            else:
                lines.append(t)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = []
                for c in row.cells:
                    ct = c.text.strip().replace("\n", " ")
                    if ct and ct not in cells:
                        cells.append(ct)
                if cells:
                    lines.append(" ، ".join(cells) + ".")
        if i % 200 == 0:
            cb(i, len(body), "قراءة المستند...")
    return "\n".join(lines), {}


def extract_pptx(path, cb):
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError("مكتبة PowerPoint غير مثبتة. نفّذ: pip install python-pptx")
    prs = Presentation(path)
    total = len(prs.slides)
    lines = []
    for i, sl in enumerate(prs.slides):
        title_shape = sl.shapes.title
        if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
            lines.append("# " + title_shape.text_frame.text.strip().replace("\n", " "))
        for sh in sl.shapes:
            if sh is title_shape or not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    lines.append(("- " + t) if len(t) < 120 and not t.endswith(".") else t)
        lines.append("")
        cb(i + 1, total, f"قراءة الشريحة {i + 1} من {total}")
    return "\n".join(lines), {"pages": total}


def extract_plain(path, cb):
    size = os.path.getsize(path)
    with open(path, "rb") as f:
        raw = f.read(MAX_TEXT_FILE_READ)
    cb(1, 1, "قراءة النص...")
    text = _decode(raw)
    text = re.sub(r"^#{1,6}\s+", "# ", text, flags=re.M)
    return text, {"truncated": size > MAX_TEXT_FILE_READ}


def extract_text(path, ext, cb, ocr=True):
    if ext == ".pdf":
        return extract_pdf(path, cb, ocr)
    if ext == ".docx":
        return extract_docx(path, cb)
    if ext == ".pptx":
        return extract_pptx(path, cb)
    return extract_plain(path, cb)


def clean_title(name):
    base = os.path.splitext(os.path.basename(name))[0]
    base = re.sub(r"[_\-]+", " ", base).strip()
    return base or "عرض"


# --------------------------------------------------------------------------
# الرفع على دفعات + المهام الخلفية
# --------------------------------------------------------------------------
UPLOADS, UPLOADS_LOCK = {}, threading.Lock()
JOBS, JOBS_LOCK = {}, threading.Lock()


def purge_old():
    now = time.time()
    with UPLOADS_LOCK:
        for uid_ in [k for k, v in UPLOADS.items() if now - v["ts"] > 3600]:
            shutil.rmtree(os.path.dirname(UPLOADS[uid_]["path"]), ignore_errors=True)
            del UPLOADS[uid_]
    with JOBS_LOCK:
        for jid in [k for k, v in JOBS.items() if now - v["ts"] > 7200]:
            del JOBS[jid]
    # بقايا من تشغيلات قديمة
    try:
        for name in os.listdir(UPLOAD_DIR):
            p = os.path.join(UPLOAD_DIR, name)
            if name not in UPLOADS and now - os.path.getmtime(p) > 6 * 3600:
                shutil.rmtree(p, ignore_errors=True)
    except OSError:
        pass


def job_update(jid, **kw):
    with JOBS_LOCK:
        if jid in JOBS:
            JOBS[jid].update(kw)


def run_convert(jid, up, target, use_ai, use_ocr=True):
    try:
        def read_cb(done, total, msg):
            job_update(jid, progress=int(3 + 55 * done / max(total, 1)), message=msg)

        job_update(jid, status="reading", progress=3, message="جاري قراءة الملف...")
        text, meta = extract_text(up["path"], up["ext"], read_cb, use_ocr)
        if meta.get("warn"):
            job_update(jid, warn=meta["warn"])
        if len(text.strip()) < 150:
            if meta.get("scanned") and not meta.get("ocr"):
                if not ocr_available():
                    raise RuntimeError("الملف ممسوح ضوئياً (صور). لقراءته ثبّت: sudo apt install tesseract-ocr tesseract-ocr-ara poppler-utils")
                raise RuntimeError("الملف ممسوح ضوئياً، فعّل خيار قراءة الصفحات الممسوحة (OCR) وحاول تاني.")
            raise RuntimeError("ما لقيتش نص كافي في الملف (حتى بعد OCR). جودة المسح ممكن تكون ضعيفة.")

        def sum_cb(frac, msg):
            job_update(jid, status="summarizing", progress=int(60 + 38 * frac), message=msg)

        deck = build_deck(text, clean_title(up["name"]), target, use_ai, sum_cb, meta)
        job_update(jid, status="done", progress=100, message="تم تجهيز السلايدات ✅", result=deck)
    except Exception as e:  # noqa
        traceback.print_exc()
        job_update(jid, status="error", error=str(e) or e.__class__.__name__, message="فشلت المعالجة")


# --------------------------------------------------------------------------
# تصدير PowerPoint
# --------------------------------------------------------------------------
def build_pptx(title, slides):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    from lxml import etree

    BG, SKY, INK, MUT = RGBColor(0x0F, 0x17, 0x2A), RGBColor(0x38, 0xBD, 0xF8), RGBColor(0xF8, 0xFA, 0xFC), RGBColor(0x94, 0xA3, 0xB8)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def style(run, size, color, bold=False):
        f = run.font
        f.size, f.bold = Pt(size), bold
        f.color.rgb = color
        f.name = "Arial"
        rpr = run._r.get_or_add_rPr()
        cs = etree.SubElement(rpr, qn("a:cs"))
        cs.set("typeface", "Arial")

    def add_text(sl, x, y, w, h, paras, size, color, bold=False, align=None, anchor=None, space=10):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = anchor
        for i, t in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            rtl = bool(AR_CHARS.search(t))
            if rtl:
                p._p.get_or_add_pPr().set("rtl", "1")
            p.alignment = align if align is not None else (PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)
            p.space_after = Pt(space)
            r = p.add_run()
            r.text = t
            style(r, size, color, bold)
        return tb

    total = len(slides)
    for i, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = BG
        bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(13.333 - 0.18), 0, Inches(0.18), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = SKY
        bar.line.fill.background()
        bullets = [b for b in s.get("bullets", []) if b]
        if s.get("type") == "title":
            add_text(sl, 1, 2.3, 11.3, 2, [s.get("title", title)], 44, SKY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            if bullets:
                add_text(sl, 1, 4.4, 11.3, 1.2, [" · ".join(bullets)], 20, MUT, False, PP_ALIGN.CENTER)
        else:
            tl = s.get("title", "")
            add_text(sl, 0.8, 0.45, 11.7, 1.2, [tl], 36, SKY, True, None, MSO_ANCHOR.MIDDLE)
            ln = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.03))
            ln.fill.solid()
            ln.fill.fore_color.rgb = RGBColor(0x33, 0x41, 0x55)
            ln.line.fill.background()
            n_chars = sum(len(b) for b in bullets)
            size = 16 if n_chars > 950 else 18 if n_chars > 700 else 20 if n_chars > 480 else 24
            add_text(sl, 0.8, 1.95, 11.7, 5.0, ["• " + b for b in bullets], size, INK, False, None, None, 12)
        add_text(sl, 0.4, 6.95, 3, 0.4, [f"{i + 1} / {total}"], 12, MUT, False, PP_ALIGN.LEFT)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def clean_slides_input(slides):
    out = []
    for s in (slides or [])[:500]:
        if isinstance(s, dict):
            out.append({"type": "title" if s.get("type") == "title" else "content",
                        "title": str(s.get("title", ""))[:300],
                        "bullets": [str(b)[:600] for b in (s.get("bullets") or [])[:40]]})
    return out


# --------------------------------------------------------------------------
# المسارات (Routes)
# --------------------------------------------------------------------------
@app.route("/")
def index():
    return Response(TEMPLATE, mimetype="text/html")


@app.route("/api/config")
def api_config():
    return jsonify({
        "ai": ai_available(),
        "pdf": have("pypdf") or have("PyPDF2"),
        "docx": have("docx"),
        "pptx": have("pptx"),
        "ocr": ocr_available(),
        "ocr_ara": ocr_available() and ocr_langs()[1],
        "max_bytes": MAX_FILE_BYTES,
    })


@app.route("/api/data", methods=["GET", "POST"])
def api_data():
    if request.method == "POST":
        body = request.get_json(silent=True)
        if not isinstance(body, dict):
            return jsonify({"status": "error", "error": "بيانات غير صالحة"}), 400
        with DATA_LOCK:
            d = load_data()
            if "events" in body:
                d["events"] = clean_events(body["events"])
            if "deck" in body:
                d["deck"] = clean_deck(body["deck"])
            if "study" in body:
                d["study"] = clean_study(body["study"])
            if "notes" in body and isinstance(body["notes"], str):
                d["notes"] = body["notes"][:100000]
            save_data(d)
        return jsonify({"status": "success"})
    return jsonify(load_data())


@app.route("/api/upload/start", methods=["POST"])
def upload_start():
    purge_old()
    body = request.get_json(silent=True) or {}
    name = str(body.get("name", "")).strip()
    try:
        size = int(body.get("size", 0))
    except (TypeError, ValueError):
        size = 0
    ext = os.path.splitext(name)[1].lower()
    if ext not in ALLOWED_EXT:
        return jsonify({"error": "نوع الملف غير مدعوم. المدعوم: PDF, DOCX, PPTX, TXT, MD, CSV"}), 400
    if size <= 0:
        return jsonify({"error": "الملف فارغ"}), 400
    if size > MAX_FILE_BYTES:
        return jsonify({"error": "الملف أكبر من 2GB"}), 413
    free = shutil.disk_usage(UPLOAD_DIR).free
    if size * 1.1 > free:
        return jsonify({"error": "مساحة القرص غير كافية لرفع الملف"}), 507
    uid_ = uuid.uuid4().hex
    folder = os.path.join(UPLOAD_DIR, uid_)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, "data" + ext)
    open(path, "wb").close()
    with UPLOADS_LOCK:
        UPLOADS[uid_] = {"name": name, "ext": ext, "size": size, "path": path,
                         "received": 0, "next": 0, "ts": time.time(), "finished": False}
    return jsonify({"id": uid_})


@app.route("/api/upload/chunk", methods=["POST"])
def upload_chunk():
    uid_ = request.args.get("id", "")
    try:
        idx = int(request.args.get("index", "-1"))
    except ValueError:
        idx = -1
    with UPLOADS_LOCK:
        up = UPLOADS.get(uid_)
    if not up:
        return jsonify({"error": "جلسة الرفع انتهت، اختر الملف تاني"}), 404
    if idx != up["next"]:
        return jsonify({"error": "ترتيب الأجزاء غير صحيح", "expected": up["next"]}), 409
    data = request.get_data()
    if up["received"] + len(data) > up["size"]:
        return jsonify({"error": "حجم البيانات أكبر من المعلن"}), 413
    with open(up["path"], "ab") as f:
        f.write(data)
    up["received"] += len(data)
    up["next"] += 1
    up["ts"] = time.time()
    return jsonify({"received": up["received"]})


@app.route("/api/upload/finish", methods=["POST"])
def upload_finish():
    body = request.get_json(silent=True) or {}
    with UPLOADS_LOCK:
        up = UPLOADS.get(str(body.get("id", "")))
    if not up:
        return jsonify({"error": "جلسة الرفع غير موجودة"}), 404
    if up["received"] != up["size"]:
        return jsonify({"error": f"الرفع ناقص ({up['received']} من {up['size']} بايت)"}), 400
    up["finished"] = True
    return jsonify({"status": "ok", "name": up["name"], "size": up["size"]})


@app.route("/api/upload/discard", methods=["POST"])
def upload_discard():
    body = request.get_json(silent=True) or {}
    with UPLOADS_LOCK:
        up = UPLOADS.pop(str(body.get("id", "")), None)
    if up:
        shutil.rmtree(os.path.dirname(up["path"]), ignore_errors=True)
    return jsonify({"status": "ok"})


@app.route("/api/convert", methods=["POST"])
def api_convert():
    body = request.get_json(silent=True) or {}
    with UPLOADS_LOCK:
        up = UPLOADS.get(str(body.get("id", "")))
    if not up or not up["finished"]:
        return jsonify({"error": "الملف لسه ما اتحملش بالكامل"}), 400
    try:
        target = int(body.get("slides") or 0)
    except (TypeError, ValueError):
        target = 0
    jid = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[jid] = {"status": "queued", "progress": 1, "message": "في الانتظار...", "ts": time.time()}
    threading.Thread(target=run_convert, args=(jid, up, target, bool(body.get("ai")), body.get("ocr", True) is not False), daemon=True).start()
    return jsonify({"job": jid})


@app.route("/api/job/<jid>")
def api_job(jid):
    with JOBS_LOCK:
        j = JOBS.get(jid)
        if not j:
            return jsonify({"error": "المهمة غير موجودة"}), 404
        out = {k: v for k, v in j.items() if k != "ts"}
    return jsonify(out)


@app.route("/api/slides/from_text", methods=["POST"])
def slides_from_text():
    body = request.get_json(silent=True) or {}
    title = str(body.get("title", "")).strip()[:200] or "عرض"
    text = str(body.get("text", ""))[:2_000_000]
    if len(text.strip()) < 100:
        return jsonify({"error": "النص قصير جداً للتلخيص"}), 400
    try:
        target = int(body.get("slides") or 0)
    except (TypeError, ValueError):
        target = 0
    try:
        deck = build_deck(text, title, target, False, lambda f, m: None, {})
    except Exception as e:  # noqa
        return jsonify({"error": str(e)}), 400
    return jsonify(deck)


@app.route("/api/export/pptx", methods=["POST"])
def export_pptx():
    if not have("pptx"):
        return jsonify({"error": "مكتبة PowerPoint غير مثبتة. نفّذ: pip install python-pptx"}), 501
    body = request.get_json(silent=True) or {}
    slides = clean_slides_input(body.get("slides"))
    if not slides:
        return jsonify({"error": "لا توجد سلايدات للتصدير"}), 400
    title = str(body.get("title", "عرض"))[:100]
    try:
        buf = build_pptx(title, slides)
    except Exception as e:  # noqa
        traceback.print_exc()
        return jsonify({"error": "فشل إنشاء ملف PowerPoint: " + str(e)}), 500
    return send_file(buf, as_attachment=True, download_name="slides.pptx",
                     mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")


@app.errorhandler(413)
def too_large(_e):
    return jsonify({"error": "الطلب أكبر من المسموح"}), 413


# --------------------------------------------------------------------------
# الواجهة
# --------------------------------------------------------------------------
TEMPLATE = r"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Abdullah Master Workspace - منظم المهام والسلايدات</title>
<link href="https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700;900&display=swap" rel="stylesheet">
<style>
:root{--bg:#0b1226;--card:#1a2338;--card2:#111a2e;--line:#2b3850;--ink:#f1f5f9;--mut:#94a3b8;--sky:#38bdf8;--blue:#2563eb;--amber:#f59e0b;--emer:#34d399;--rose:#fb7185}
*{box-sizing:border-box}
html,body{margin:0}
body{background:var(--bg);color:var(--ink);font-family:'Cairo','Segoe UI',Tahoma,Arial,sans-serif;min-height:100vh;display:flex;flex-direction:column}
button{font-family:inherit;cursor:pointer;border:0}
input,select,textarea{font-family:inherit;font-size:15px}
.hidden{display:none!important}
body.no-scroll{overflow:hidden}

/* الهيدر */
.topbar{background:#111a2e;border-bottom:1px solid var(--line);padding:14px 28px;display:flex;align-items:center;gap:18px;position:sticky;top:0;z-index:40;flex-wrap:wrap}
.brand{display:flex;align-items:center;gap:10px}
.brand .logo{font-size:30px}
.brand h1{margin:0;font-size:20px;font-weight:800;background:linear-gradient(90deg,#38bdf8,#3b82f6);-webkit-background-clip:text;background-clip:text;color:transparent}
.tabs{display:flex;gap:8px;margin-inline-start:auto}
.tab-btn{padding:9px 16px;border-radius:10px;font-weight:700;font-size:14px;background:#161f35;color:#cbd5e1;border:1px solid transparent;transition:.15s}
.tab-btn:hover{color:#fff;border-color:var(--line)}
.tab-btn.active{background:var(--blue);color:#fff}
.top-extra{display:flex;align-items:center;gap:10px;margin-inline-start:auto}
.pill{border:1px solid var(--line);border-radius:999px;padding:6px 14px;font-size:12px;color:var(--mut);background:#0f172a;white-space:nowrap}
.notif{background:rgba(245,158,11,.15);border:1px solid rgba(245,158,11,.5);color:#fbbf24;border-radius:999px;padding:6px 12px;font-size:12px;font-weight:700}
.notif.on{background:rgba(52,211,153,.15);border-color:rgba(52,211,153,.5);color:var(--emer)}

main{flex:1;width:100%;max-width:1500px;margin:0 auto;padding:30px 28px 40px}
footer{border-top:1px solid var(--line);background:#111a2e;text-align:center;font-size:12px;color:var(--mut);padding:16px}

.card{background:var(--card);border:1px solid var(--line);border-radius:18px;padding:22px;box-shadow:0 10px 30px rgba(0,0,0,.25)}
.btn{padding:11px 18px;border-radius:12px;font-weight:700;font-size:15px;transition:.15s;color:#fff}
.btn:hover:not(:disabled){filter:brightness(1.1)}
.btn:disabled{opacity:.55;cursor:not-allowed}
.btn.blue{background:var(--blue)}
.btn.amber{background:#d97706}
.btn.ghost{background:#1e293b;color:#cbd5e1;border:1px solid var(--line)}
.btn.sky{background:linear-gradient(90deg,#38bdf8,#2563eb);color:#06101f}
.btn.green{background:linear-gradient(90deg,#34d399,#0d9488);color:#04120f}
.field label{display:block;font-size:13px;color:var(--mut);margin-bottom:7px;font-weight:600}
.field input,.field select,.field textarea{width:100%;background:#0e1629;border:1px solid var(--line);border-radius:12px;padding:12px 14px;color:var(--ink);outline:none;transition:.15s}
.field input:focus,.field select:focus,.field textarea:focus{border-color:var(--sky)}
.field textarea{resize:vertical;min-height:170px;line-height:1.8}
::placeholder{color:#64748b}

/* الجدول */
.cal-wrap{max-width:1000px;margin:0 auto;display:flex;flex-direction:column;gap:18px}
.hero{text-align:center;margin-bottom:4px}
.hero h2{margin:6px 0 6px;font-size:38px;font-weight:900;background:linear-gradient(90deg,#38bdf8,#3b82f6);-webkit-background-clip:text;background-clip:text;color:transparent}
.hero p{margin:0;color:var(--mut)}
.add-grid{display:grid;grid-template-columns:2fr 1.1fr 1.5fr 1.1fr auto;gap:14px;align-items:end}
.navbar{display:flex;align-items:center;justify-content:space-between;gap:12px;flex-wrap:wrap;padding:14px 18px}
.seg{display:flex;gap:8px}
.seg button{padding:9px 16px;border-radius:10px;background:#0e1629;color:#cbd5e1;border:1px solid var(--line);font-weight:700;font-size:14px}
.seg button.active{background:var(--sky);color:#06101f;border-color:var(--sky)}
.pager{display:flex;align-items:center;gap:12px;flex-wrap:wrap;justify-content:center;max-width:100%}
.round{width:38px;height:38px;border-radius:50%;background:#0e1629;border:1px solid var(--line);color:#e2e8f0;font-size:16px}
.round:hover{border-color:var(--sky)}
.period{font-weight:800;color:var(--sky);min-width:0;text-align:center;font-size:15px}
.today-btn{background:#0e1629;border:1px solid var(--line);color:#cbd5e1;border-radius:10px;padding:7px 12px;font-weight:700;font-size:13px}
.today-btn:hover{border-color:var(--sky)}
.days{display:flex;flex-direction:column;gap:14px}
.day{background:var(--card);border:1px solid var(--line);border-radius:18px;padding:18px 22px;transition:.15s}
.day:hover{border-color:#3b5b85}
.day.today{border-color:var(--sky);box-shadow:0 0 0 1px rgba(56,189,248,.35)}
.day-head{display:flex;justify-content:space-between;align-items:center;border-bottom:1px solid var(--line);padding-bottom:12px;margin-bottom:12px}
.day-name{font-size:20px;font-weight:800;display:flex;align-items:center;gap:10px}
.badge{font-size:11px;background:var(--sky);color:#06101f;border-radius:999px;padding:2px 10px;font-weight:800}
.date-pill{font-size:12px;font-weight:700;color:var(--sky);background:rgba(14,165,233,.12);border:1px solid rgba(56,189,248,.25);border-radius:999px;padding:4px 12px}
.empty{color:#64748b;font-style:italic;text-align:center;padding:8px 0;margin:0;font-size:14px}
.ev{display:flex;align-items:center;gap:12px;background:#0e1629;border:1px solid var(--line);border-radius:12px;padding:10px 14px;margin-top:8px}
.ev:first-of-type{margin-top:0}
.ev input[type=checkbox]{width:18px;height:18px;accent-color:var(--emer);cursor:pointer}
.ev-main{flex:1;display:flex;flex-direction:column;gap:2px;min-width:0}
.ev-title{font-weight:700;overflow-wrap:anywhere}
.ev-time{font-size:12px;color:var(--mut)}
.ev.done .ev-title{text-decoration:line-through;color:#64748b}
.del{background:rgba(244,63,94,.12);color:#fb7185;border-radius:8px;padding:5px 12px;font-size:12px;font-weight:700}
.del:hover{background:rgba(244,63,94,.25)}
.month-grid{display:grid;grid-template-columns:repeat(7,1fr);gap:8px}
.mh{text-align:center;font-size:12px;font-weight:700;color:var(--mut);padding:4px 0}
.cell{min-height:92px;background:var(--card);border:1px solid var(--line);border-radius:12px;padding:8px;text-align:right;color:var(--ink);display:flex;flex-direction:column;gap:4px;overflow:hidden}
.cell:hover{border-color:#3b5b85}
.cell.blank{background:transparent;border-color:transparent;pointer-events:none}
.cell.today .dn{background:var(--sky);color:#06101f}
.cell.sel{border-color:var(--sky);box-shadow:0 0 0 1px rgba(56,189,248,.4)}
.dn{font-size:13px;font-weight:800;width:24px;height:24px;border-radius:50%;display:flex;align-items:center;justify-content:center}
.chip{font-size:11px;background:rgba(56,189,248,.14);border-radius:6px;padding:2px 6px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:#bae6fd}
.chip.done{opacity:.5;text-decoration:line-through}
.more{font-size:11px;color:var(--mut)}

/* السلايدات */
.slides-grid{display:grid;grid-template-columns:1fr 1fr;gap:26px;align-items:start}
.card h3.t{margin:0 0 18px;font-size:20px;font-weight:800;display:flex;align-items:center;gap:10px}
.t.amber{color:#fbbf24}.t.emer{color:var(--emer)}
.stack{display:flex;flex-direction:column;gap:16px}
.filebox{display:flex;align-items:center;gap:12px;background:#0e1629;border:1px solid var(--line);border-radius:12px;padding:10px 12px}
.browse{background:var(--blue);color:#fff;font-weight:700;padding:9px 16px;border-radius:9px;font-size:14px;cursor:pointer;white-space:nowrap}
.browse:hover{filter:brightness(1.1)}
.fname{color:var(--mut);font-size:14px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;direction:ltr;text-align:right;flex:1}
.hint{font-size:12px;color:#64748b;margin-top:6px}
.prog-head{display:flex;justify-content:space-between;font-size:12px;color:var(--mut);margin-bottom:6px;gap:12px}
.bar-bg{height:12px;background:#0e1629;border:1px solid var(--line);border-radius:999px;overflow:hidden}
.bar-fg{height:100%;width:0;background:var(--amber);transition:width .25s}
.bar-fg.ok{background:var(--emer)}.bar-fg.err{background:var(--rose)}
.convert{width:100%;padding:13px;border-radius:12px;font-weight:800;font-size:16px;color:#1b1203;background:linear-gradient(90deg,#f59e0b,#ea580c)}
.convert:disabled{background:#334155;color:#94a3b8}
.row2{display:grid;grid-template-columns:1fr 1fr;gap:12px}
.feat{background:#0e1629;border:1px solid var(--line);border-radius:12px;padding:14px 16px;margin-top:20px}
.feat b{color:var(--sky);font-size:14px}
.feat ul{margin:8px 0 0;padding-inline-start:20px;color:var(--mut);font-size:12.5px;line-height:1.9}
.chk-row{display:flex;align-items:center;gap:8px;font-size:13px;color:#cbd5e1}
.chk-row input{accent-color:var(--emer);width:16px;height:16px}

/* المحرر */
.overlay{position:fixed;inset:0;background:var(--bg);z-index:100;display:flex;flex-direction:column}
.ed-top{display:flex;gap:14px;align-items:center;padding:12px 20px;background:#111a2e;border-bottom:1px solid var(--line);flex-wrap:wrap}
.ed-title-input{flex:1;min-width:200px;background:#0e1629;border:1px solid var(--line);border-radius:10px;padding:10px 14px;color:var(--ink);font-size:17px;font-weight:800;outline:none}
.ed-title-input:focus{border-color:var(--sky)}
.ed-actions{display:flex;gap:8px;flex-wrap:wrap}
.ed-actions .btn{padding:9px 14px;font-size:14px}
.ed-body{flex:1;display:flex;min-height:0}
.ed-side{width:270px;border-inline-end:1px solid var(--line);background:#0e1629;display:flex;flex-direction:column;min-height:0}
#ed-list{flex:1;overflow-y:auto;padding:10px;display:flex;flex-direction:column;gap:6px}
.thumb{display:flex;gap:10px;align-items:center;text-align:right;background:#162036;color:#cbd5e1;border:1px solid transparent;border-radius:10px;padding:9px 10px;width:100%}
.thumb:hover{border-color:var(--line)}
.thumb.active{border-color:var(--sky);background:#1a2b47;color:#fff}
.thumb .n{background:#0b1226;border-radius:6px;min-width:26px;height:26px;display:flex;align-items:center;justify-content:center;font-size:12px;font-weight:800;color:var(--sky)}
.thumb .tt{font-size:13px;font-weight:600;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;flex:1}
.ed-side .add{margin:10px}
.ed-main{flex:1;overflow-y:auto;padding:22px;display:flex;flex-direction:column;gap:18px;align-items:center}
.ed-main>*{width:100%;max-width:960px}
.ed-tools{display:flex;gap:8px;flex-wrap:wrap}
.ed-tools .btn{padding:8px 14px;font-size:13px}
.slide-box{position:relative;width:100%;aspect-ratio:16/9;overflow:hidden;border-radius:14px;border:1px solid var(--line);background:#0f172a}

/* شكل السلايد (1280×720) */
.slide{position:absolute;top:0;left:0;width:1280px;height:720px;transform-origin:0 0;direction:rtl;background:linear-gradient(135deg,#0f172a 0%,#1b2740 100%);color:#f8fafc;font-family:'Cairo','Segoe UI',Tahoma,Arial,sans-serif;overflow:hidden}
.slide .accent{position:absolute;top:0;right:0;bottom:0;width:16px;background:linear-gradient(180deg,#38bdf8,#2563eb)}
.slide h2{position:absolute;top:56px;right:90px;left:90px;margin:0;font-size:54px;line-height:1.25;font-weight:900;color:#38bdf8;padding-bottom:20px;border-bottom:3px solid #334155;overflow:hidden;max-height:130px}
.slide ul{position:absolute;top:210px;right:90px;left:90px;bottom:70px;margin:0;padding:0;list-style:none;overflow:hidden}
.slide li{position:relative;padding-inline-start:38px;margin-bottom:16px;line-height:1.55;color:#e2e8f0}
.slide li::before{content:"";position:absolute;inset-inline-start:0;top:.62em;width:14px;height:14px;border-radius:50%;background:#38bdf8}
.slide.lg li{font-size:36px}.slide.md li{font-size:31px}.slide.sm li{font-size:27px}.slide.xs li{font-size:23px;margin-bottom:12px}
.slide .num{position:absolute;bottom:26px;left:60px;font-size:20px;color:#64748b}
.slide .brand-s{position:absolute;bottom:26px;right:90px;font-size:18px;color:#475569;direction:ltr}
.slide.cover{display:flex;flex-direction:column;align-items:center;justify-content:center;text-align:center;padding:0 140px;background:radial-gradient(circle at 70% 20%,#1d3a63 0%,#0f172a 60%)}
.slide.cover .em{font-size:96px;margin-bottom:18px}
.slide.cover h1{margin:0;font-size:68px;line-height:1.3;font-weight:900;color:#38bdf8;max-height:300px;overflow:hidden}
.slide.cover p{margin:26px 0 0;font-size:28px;color:#94a3b8}

/* وضع العرض */
.present{position:fixed;inset:0;background:#000;z-index:200;display:flex;align-items:center;justify-content:center}
.present .slide-box{border:0;border-radius:0;width:auto}
.pres-ui{position:fixed;bottom:16px;left:50%;transform:translateX(-50%);display:flex;gap:10px;align-items:center;background:rgba(15,23,42,.85);border:1px solid var(--line);border-radius:999px;padding:6px 14px;opacity:.25;transition:.2s}
.pres-ui:hover{opacity:1}
.pres-ui button{background:#1e293b;color:#fff;border-radius:50%;width:36px;height:36px;font-size:16px}
.pres-ui span{color:#cbd5e1;font-size:14px;min-width:60px;text-align:center}

/* قايمة المشاريع ⋮ */
.cur-proj{font-size:13px;font-weight:700;color:#bae6fd;background:rgba(56,189,248,.12);border:1px solid rgba(56,189,248,.3);border-radius:999px;padding:5px 14px;white-space:nowrap}
.proj-menu{position:relative}
.kebab{width:42px;height:42px;border-radius:12px;background:#161f35;border:1px solid var(--line);color:#e2e8f0;font-size:26px;line-height:1;display:flex;align-items:center;justify-content:center;padding:0 0 4px}
.kebab:hover,.kebab[aria-expanded=true]{border-color:var(--sky);color:#fff;background:#1a2b47}
.menu{position:absolute;top:calc(100% + 8px);inset-inline-end:0;min-width:270px;max-width:92vw;background:#111a2e;border:1px solid var(--line);border-radius:14px;padding:8px;box-shadow:0 18px 40px rgba(0,0,0,.55);z-index:60;display:flex;flex-direction:column;gap:4px;animation:tin .15s}
.menu-title{font-size:12px;color:var(--mut);padding:6px 10px 8px;font-weight:700}
.menu button{display:flex;align-items:center;gap:10px;text-align:right;width:100%;padding:11px 12px;border-radius:10px;background:transparent;color:#e2e8f0;font-weight:700;font-size:14px}
.menu button:hover{background:#1a2b47}
.menu button.active{background:var(--blue);color:#fff}

/* متابعة المنهج */
.study-wrap{max-width:1000px;margin:0 auto;display:flex;flex-direction:column;gap:20px}
.st-row{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap;margin-bottom:14px}
.st-row h3{margin:0 0 4px;font-size:18px;font-weight:800}
.st-row small{color:var(--mut);font-size:12px}
.st-big{font-size:36px;font-weight:900;color:var(--sky)}
.st-bar-bg{height:16px;background:#0e1629;border:1px solid var(--line);border-radius:999px;padding:2px;overflow:hidden}
.st-bar-fg{height:100%;width:0;border-radius:999px;background:linear-gradient(90deg,#06b6d4,#3b82f6,#6366f1);transition:width .6s ease}
.chips{display:flex;gap:8px;overflow-x:auto;padding-bottom:12px;border-bottom:1px solid var(--line)}
.chip-btn{display:flex;align-items:center;gap:8px;padding:10px 16px;border-radius:12px;font-weight:700;font-size:14px;white-space:nowrap;background:#111a2e;border:1px solid var(--line);color:#cbd5e1}
.chip-btn:hover{background:#1a2338;color:#fff}
.chip-btn.active{background:linear-gradient(135deg,#0284c7,#2563eb);border-color:rgba(56,189,248,.5);color:#fff;box-shadow:0 4px 20px -2px rgba(14,165,233,.4)}
.st-badge{font-size:12px;padding:1px 9px;border-radius:999px;background:#1e293b;color:#94a3b8}
.chip-btn.active .st-badge{background:rgba(255,255,255,.2);color:#fff}
.subj-head{display:flex;justify-content:space-between;align-items:center;gap:14px;flex-wrap:wrap;padding-bottom:18px;margin-bottom:18px;border-bottom:1px solid var(--line)}
.subj-title{display:flex;align-items:center;gap:12px}
.subj-ico{width:48px;height:48px;border-radius:12px;background:#1e293b;border:1px solid var(--line);display:flex;align-items:center;justify-content:center;font-size:24px}
.subj-title h2{margin:0;font-size:24px;font-weight:900}
.subj-title p{margin:2px 0 0;font-size:12px;color:var(--mut)}
.subj-pct{background:#1e293b;border:1px solid var(--line);border-radius:12px;padding:8px 16px}
.subj-pct small{display:block;font-size:12px;color:var(--mut)}
.subj-pct b{font-size:18px;color:var(--sky);font-weight:900}
.grp{background:#0e1629;border:1px solid var(--line);border-radius:14px;padding:14px;margin-bottom:14px}
.grp h4{margin:0 0 10px;padding-bottom:8px;border-bottom:1px solid var(--line);font-size:16px;font-weight:800;color:var(--sky)}
.grp.custom h4{color:#fbbf24}
.lesson{display:flex;align-items:center;gap:12px;padding:10px 12px;border-radius:12px;border:1px solid var(--line);background:rgba(30,41,59,.5);margin-top:8px;cursor:pointer;user-select:none;transition:.15s}
.lesson:hover{background:rgba(30,41,59,.9)}
.lesson.done{opacity:.65;background:rgba(15,23,42,.5)}
.lesson.done .lt{text-decoration:line-through;color:#64748b}
.lesson input{width:19px;height:19px;accent-color:#06b6d4;cursor:pointer;flex:none}
.lt{font-size:14px;font-weight:600;flex:1;overflow-wrap:anywhere}
.lesson .del{margin-inline-start:auto}
.add-lesson{padding-top:16px;border-top:1px solid var(--line)}
.add-lesson h5{margin:0 0 8px;font-size:12px;color:var(--mut);font-weight:700}
.add-lesson .r{display:flex;gap:8px;flex-wrap:wrap}
.add-lesson input{flex:1;min-width:180px;background:#1e293b;border:1px solid var(--line);border-radius:12px;padding:10px 14px;color:var(--ink);outline:none}
.add-lesson input:focus{border-color:var(--sky)}

/* التوست */
#toasts{position:fixed;top:78px;left:50%;transform:translateX(-50%);display:flex;flex-direction:column;gap:8px;z-index:500;align-items:center;pointer-events:none}
.toast{background:#1e293b;border:1px solid var(--line);border-radius:12px;padding:11px 20px;font-weight:700;font-size:14px;box-shadow:0 10px 30px rgba(0,0,0,.45);animation:tin .25s;max-width:90vw}
.toast.ok{border-color:rgba(52,211,153,.6);color:#a7f3d0}.toast.warn{border-color:rgba(245,158,11,.6);color:#fde68a}.toast.err{border-color:rgba(251,113,133,.6);color:#fecdd3}
@keyframes tin{from{opacity:0;transform:translateY(-8px)}to{opacity:1;transform:none}}

#print-root{display:none}
@media print{
  @page{size:13.333in 7.5in;margin:0}
  html,body{background:#0f172a!important;-webkit-print-color-adjust:exact;print-color-adjust:exact}
  body>*:not(#print-root){display:none!important}
  #print-root{display:block}
  .print-page{width:1280px;height:720px;position:relative;overflow:hidden;break-after:page;page-break-after:always}
  .print-page:last-child{break-after:auto;page-break-after:auto}
  .print-page .slide{transform:none!important}
}
@media (max-width:900px){
  .slides-grid{grid-template-columns:1fr}
  .add-grid{grid-template-columns:1fr 1fr}
  .ed-side{width:190px}
  main{padding:20px 14px}
  .topbar{padding:12px 14px}
  .hero h2{font-size:28px}
}
</style>
</head>
<body>

<header class="topbar">
  <div class="brand"><span class="logo">🐉</span><h1>Abdullah Master Workspace</h1></div>
  <span class="cur-proj" id="cur-proj"></span>
  <div class="top-extra">
    <button class="notif" id="notif-btn" onclick="requestNotif()">🔔 تفعيل الإشعارات</button>
    <span class="pill">سنة تانية - هندسة وحاسبات</span>
    <div class="proj-menu">
      <button class="kebab" id="kebab-btn" aria-haspopup="true" aria-expanded="false" aria-label="تبديل المشروع" title="تبديل المشروع">⋮</button>
      <div class="menu hidden" id="proj-menu" role="menu">
        <div class="menu-title">انتقل إلى مشروع</div>
        <button role="menuitem" data-proj="calendar">📅 الجداول والمواعيد والمهام</button>
        <button role="menuitem" data-proj="slides">📊 مولد السلايدات والمحرر</button>
        <button role="menuitem" data-proj="study">🎓 متابعة المنهج والمذاكرة</button>
      </div>
    </div>
  </div>
</header>

<main>
<!-- ============ الجدول ============ -->
<section id="tab-calendar" class="cal-wrap">
  <div class="hero">
    <h2>منظم الأحداث والجداول الذكي</h2>
    <p>إدارة الأسبوع والشهر بكل سلاسة واحترافية</p>
  </div>

  <div class="card add-grid">
    <div class="field"><label>الحدث / المهمة</label><input type="text" id="event-input" maxlength="300" placeholder="مثال: امتحان تخصص / تمرين جيم"></div>
    <div class="field"><label>الميعاد (الوقت)</label><input type="time" id="time-input" value="08:00"></div>
    <div class="field"><label>اليوم</label><select id="day-select"></select></div>
    <div class="field"><label>التنبيه</label>
      <select id="remind-select">
        <option value="0">في الميعاد</option><option value="5">قبلها بـ 5 دقائق</option>
        <option value="10">قبلها بـ 10 دقائق</option><option value="30">قبلها بـ 30 دقيقة</option>
        <option value="60">قبلها بساعة</option>
      </select></div>
    <button class="btn sky" id="add-btn" style="height:47px">إضافة الحدث</button>
  </div>

  <div class="card navbar">
    <div class="seg">
      <button id="view-week" class="active">عرض الأسبوع</button>
      <button id="view-month">عرض الشهر</button>
    </div>
    <div class="pager">
      <button class="round" id="prev-btn" title="السابق">→</button>
      <span class="period" id="period-label"></span>
      <button class="round" id="next-btn" title="التالي">←</button>
      <button class="today-btn" id="today-btn">اليوم</button>
    </div>
  </div>

  <div id="cal-body"></div>
</section>

<!-- ============ السلايدات ============ -->
<section id="tab-slides" class="hidden">
  <div class="slides-grid">
    <!-- رفع ملف (يظهر يمين في الـ RTL) -->
    <div class="card" style="display:flex;flex-direction:column">
      <h3 class="t amber"><span>📂</span> رفع ملف ملخص أو ملاحظات</h3>
      <div class="stack">
        <div class="field">
          <label>اختر ملف نصي أو ملخص أو كتاب كامل:</label>
          <div class="filebox">
            <label class="browse">Browse...<input type="file" id="upload-file" class="hidden" accept=".pdf,.docx,.pptx,.txt,.md,.csv"></label>
            <span class="fname" id="file-name">No file selected.</span>
          </div>
          <div class="hint">PDF · Word · PowerPoint · TXT · MD — حتى 2GB، بيترفع على دفعات.</div>
        </div>

        <div class="row2">
          <div class="field"><label>عدد السلايدات</label>
            <select id="slides-count">
              <option value="0">تلقائي</option><option value="10">10</option><option value="15">15</option>
              <option value="20">20</option><option value="30">30</option><option value="50">50</option>
            </select></div>
          <div class="field" id="ai-wrap" style="display:flex;align-items:flex-end">
            <label class="chk-row"><input type="checkbox" id="ai-check"> ✨ تحسين بـ Claude</label>
          </div>
        </div>

        <div id="ocr-wrap">
          <label class="chk-row" id="ocr-row"><input type="checkbox" id="ocr-check" checked> 🔍 قراءة الصفحات الممسوحة ضوئياً (OCR) تلقائياً</label>
          <div class="hint hidden" id="ocr-hint">لقراءة الكتب الممسوحة ثبّت: <span dir="ltr">sudo apt install tesseract-ocr tesseract-ocr-ara poppler-utils</span></div>
        </div>

        <div id="progress-box" class="hidden">
          <div class="prog-head"><span id="prog-text">جاري رفع الملف...</span><span id="prog-pct">0%</span></div>
          <div class="bar-bg"><div class="bar-fg" id="prog-bar"></div></div>
        </div>

        <button class="convert" id="convert-btn" disabled>تحويل الملف إلى سلايدات 📤</button>
        <button class="btn ghost hidden" id="last-deck-btn">📂 فتح آخر عرض محفوظ</button>
      </div>

      <div class="feat">
        <b>💡 ميزات النظام:</b>
        <ul>
          <li>الزر بيشتغل بعد اكتمال الرفع الفعلي، والمعالجة بتتم في الخلفية مع شريط تقدم حقيقي.</li>
          <li>تلخيص ذكي محلي: بيقسم الكتاب لأقسام ويطلع أهم الجمل في كل قسم (عربي وإنجليزي).</li>
          <li>الكتب الممسوحة ضوئياً: بيتعمل OCR للصفحات اللي مفيهاش نص (محتاج tesseract).</li>
          <li>محرر كامل + عرض تقديمي + تصدير PDF و PowerPoint.</li>
          <li>لتفعيل تلخيص Claude: ضيف المتغير ANTHROPIC_API_KEY قبل تشغيل البرنامج.</li>
        </ul>
      </div>
    </div>

    <!-- كتابة المحتوى -->
    <div class="card">
      <h3 class="t emer"><span>📝</span> كتابة محتوى السلايدات أو الدروس</h3>
      <div class="stack">
        <div class="field"><label>عنوان الدرس أو السلايد الرئيسي:</label>
          <input type="text" id="slide-title" maxlength="200" placeholder="مثال: هندسة البرمجيات الحديثة"></div>
        <div class="field"><label>النقاط (كل نقطة في سطر مستقل):</label>
          <textarea id="slide-points" rows="8" placeholder="النقطة الأولى&#10;النقطة الثانية&#10;النقطة الثالثة"></textarea>
          <div class="hint">ابدأ سطر بـ <b dir="ltr"># </b> لعمل سلايد جديد بعنوان، وإلا بتتوزع كل 6 نقاط في سلايد. لو لصقت نص طويل (أكثر من 1200 حرف) هيتلخص تلقائياً.</div>
        </div>
        <label class="chk-row"><input type="checkbox" id="smart-check" checked> تلخيص ذكي تلقائي للنص الطويل</label>
        <button class="btn blue" id="gen-btn" style="width:100%;padding:13px;font-size:16px">🚀 توليد السلايدات وتصديرها (PDF / PPTX)</button>
      </div>
    </div>
  </div>
</section>

<!-- ============ متابعة المنهج ============ -->
<section id="tab-study" class="study-wrap hidden">
  <div class="hero">
    <h2>منصة متابعة المنهج - عبد الله 🚀</h2>
    <p>تابع تقدمك الدراسي بكل دقة، نظم جدولك، وحقق أهدافك خطوة بخطوة.</p>
  </div>
  <div class="card">
    <div class="st-row">
      <div><h3>📈 التقدم الإجمالي للمنهج الكلي</h3><small>إجمالي جميع المواد والدروس المستكملة</small></div>
      <div><span class="st-big" id="st-pct">0%</span> <small id="st-count">0 من 0 درس</small></div>
    </div>
    <div class="st-bar-bg"><div class="st-bar-fg" id="st-bar"></div></div>
  </div>
  <div class="chips" id="st-tabs"></div>
  <div id="st-views" class="stack" style="gap:22px"></div>
</section>
</main>

<footer>Abdullah Master Nexus OS &copy; 2026 - Designed for Ultimate Productivity</footer>

<!-- ============ المحرر ============ -->
<div id="editor" class="overlay hidden">
  <div class="ed-top">
    <input id="ed-title" class="ed-title-input" placeholder="عنوان العرض" maxlength="200">
    <div class="ed-actions">
      <button class="btn ghost" id="btn-present">▶ عرض تقديمي</button>
      <button class="btn blue" id="btn-pdf">📄 تصدير PDF</button>
      <button class="btn amber" id="btn-pptx">📑 تصدير PowerPoint</button>
      <button class="btn ghost" id="btn-ed-close">✕ إغلاق</button>
    </div>
  </div>
  <div class="ed-body">
    <aside class="ed-side">
      <div id="ed-list"></div>
      <button class="btn green add" id="btn-add-slide">＋ سلايد جديد</button>
    </aside>
    <section class="ed-main">
      <div class="slide-box" id="ed-preview"></div>
      <div class="card stack">
        <div class="field"><label>عنوان السلايد</label><input type="text" id="ed-slide-title" maxlength="300"></div>
        <div class="field"><label>النقاط (كل نقطة في سطر)</label><textarea id="ed-slide-bullets" rows="7"></textarea></div>
        <div class="ed-tools">
          <button class="btn ghost" id="btn-up">↑ تحريك لأعلى</button>
          <button class="btn ghost" id="btn-down">↓ تحريك لأسفل</button>
          <button class="btn ghost" id="btn-dup">⧉ تكرار</button>
          <button class="btn ghost" id="btn-del" style="color:#fb7185">🗑 حذف السلايد</button>
        </div>
      </div>
    </section>
  </div>
</div>

<div id="present" class="present hidden">
  <div class="slide-box" id="present-box"></div>
  <div class="pres-ui">
    <button id="pres-prev" title="السابق">→</button><span id="pres-count"></span>
    <button id="pres-next" title="التالي">←</button><button id="pres-exit" title="خروج">✕</button>
  </div>
</div>

<div id="print-root"></div>
<div id="toasts"></div>

<script>
/* ---------- أدوات عامة ---------- */
const $ = id => document.getElementById(id);
const esc = s => String(s == null ? '' : s).replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const sleep = ms => new Promise(r => setTimeout(r, ms));
const pad = n => String(n).padStart(2, '0');
const uid = () => Math.random().toString(36).slice(2, 10) + Date.now().toString(36).slice(-4);
let CONFIG = {ai:false, pdf:true, docx:true, pptx:true};

function toast(msg, type) {
  const t = document.createElement('div');
  t.className = 'toast ' + (type || 'ok');
  t.textContent = msg;
  $('toasts').appendChild(t);
  setTimeout(() => { t.style.opacity = '0'; t.style.transition = '.3s'; setTimeout(() => t.remove(), 300); }, 3200);
}
async function postJSON(url, body) {
  const r = await fetch(url, {method:'POST', headers:{'Content-Type':'application/json'}, body: JSON.stringify(body)});
  let j = {};
  try { j = await r.json(); } catch (e) {}
  if (!r.ok) throw new Error(j.error || ('خطأ ' + r.status));
  return j;
}
function lsGet(k) { try { return localStorage.getItem(k); } catch (e) { return null; } }
function lsSet(k, v) { try { localStorage.setItem(k, v); } catch (e) {} }

/* ---------- المشاريع + قايمة ⋮ ---------- */
const PROJECTS = [
  {id: 'calendar', icon: '📅', name: 'الجداول والمواعيد والمهام'},
  {id: 'slides', icon: '📊', name: 'مولد السلايدات والمحرر'},
  {id: 'study', icon: '🎓', name: 'متابعة المنهج والمذاكرة'}
];
function openMenu() { $('proj-menu').classList.remove('hidden'); $('kebab-btn').setAttribute('aria-expanded', 'true'); }
function closeMenu() { $('proj-menu').classList.add('hidden'); $('kebab-btn').setAttribute('aria-expanded', 'false'); }
function switchTab(name) {
  if (!PROJECTS.some(p => p.id === name)) name = 'calendar';
  PROJECTS.forEach(p => $('tab-' + p.id).classList.toggle('hidden', p.id !== name));
  const cur = PROJECTS.find(p => p.id === name);
  $('cur-proj').textContent = cur.icon + ' ' + cur.name;
  document.querySelectorAll('#proj-menu [data-proj]').forEach(b => b.classList.toggle('active', b.dataset.proj === name));
  closeMenu();
  lsSet('tab', name);
  window.scrollTo(0, 0);
}
$('kebab-btn').addEventListener('click', e => { e.stopPropagation(); $('proj-menu').classList.contains('hidden') ? openMenu() : closeMenu(); });
$('proj-menu').addEventListener('click', e => { const b = e.target.closest('[data-proj]'); if (b) switchTab(b.dataset.proj); });
document.addEventListener('click', e => { if (!e.target.closest('.proj-menu')) closeMenu(); });
document.addEventListener('keydown', e => { if (e.key === 'Escape') closeMenu(); });

/* ---------- التاريخ ---------- */
const DAYS = ['السبت','الأحد','الإثنين','الثلاثاء','الأربعاء','الخميس','الجمعة'];
const MONTHS = ['يناير','فبراير','مارس','أبريل','مايو','يونيو','يوليو','أغسطس','سبتمبر','أكتوبر','نوفمبر','ديسمبر'];
const fmt = d => d.getFullYear() + '-' + pad(d.getMonth() + 1) + '-' + pad(d.getDate());
const dayIdx = d => (d.getDay() + 1) % 7;               // السبت = 0
const addDays = (d, n) => new Date(d.getFullYear(), d.getMonth(), d.getDate() + n);
const weekStart = d => addDays(d, -dayIdx(d));
const short = d => d.getDate() + ' / ' + (d.getMonth() + 1);
function fmtTime(t) {
  const [h, m] = String(t).split(':').map(Number);
  if (isNaN(h)) return t;
  return ((h % 12) || 12) + ':' + pad(m || 0) + ' ' + (h >= 12 ? 'م' : 'ص');
}
function parseDate(s) { const [y, m, d] = s.split('-').map(Number); return new Date(y, m - 1, d); }

/* ---------- حالة الجدول ---------- */
const T = new Date();
let anchor = new Date(T.getFullYear(), T.getMonth(), T.getDate());
let view = 'week';
let selected = fmt(anchor);
let events = {};

function periodDays() {
  if (view === 'week') {
    const s = weekStart(anchor);
    return Array.from({length: 7}, (_, i) => addDays(s, i));
  }
  const n = new Date(anchor.getFullYear(), anchor.getMonth() + 1, 0).getDate();
  return Array.from({length: n}, (_, i) => new Date(anchor.getFullYear(), anchor.getMonth(), i + 1));
}

async function loadData() {
  try {
    const r = await fetch('/api/data');
    const d = await r.json();
    events = d.events || {};
    Object.values(events).forEach(list => list.forEach(e => {
      if (!e.id) e.id = uid();
      if (e.done == null) e.done = false;
      if (e.remind == null) e.remind = 0;
    }));
    if (d.deck && d.deck.slides && d.deck.slides.length) setLastDeck(d.deck);
    STUDY.progress = (d.study && d.study.progress) || {};
    STUDY.custom = (d.study && d.study.custom) || {};
  } catch (e) {
    events = {};
    toast('تعذر تحميل البيانات المحفوظة', 'err');
  }
  renderCal();
  renderStudy();
}

let saving = Promise.resolve();
function persistEvents() {
  saving = saving.then(() => postJSON('/api/data', {events}))
    .catch(() => toast('تعذر حفظ التغييرات في الملف', 'err'));
  return saving;
}

function renderCal() {
  const days = periodDays();
  if (!days.some(d => fmt(d) === selected)) {
    const t = days.find(d => fmt(d) === fmt(T)) || days[0];
    selected = fmt(t);
  }
  const sel = $('day-select');
  sel.innerHTML = days.map(d => `<option value="${fmt(d)}">${DAYS[dayIdx(d)]} (${pad(d.getDate())}-${pad(d.getMonth() + 1)}-${d.getFullYear()})</option>`).join('');
  sel.value = selected;

  $('view-week').classList.toggle('active', view === 'week');
  $('view-month').classList.toggle('active', view === 'month');
  if (view === 'week') {
    $('period-label').textContent = `الأسبوع (من ${short(days[0])} إلى ${short(days[6])})`;
    $('cal-body').innerHTML = '<div class="days">' + days.map(dayCardHTML).join('') + '</div>';
  } else {
    $('period-label').textContent = MONTHS[anchor.getMonth()] + ' ' + anchor.getFullYear();
    renderMonth(days);
  }
}

function eventsHTML(ds) {
  const list = (events[ds] || []).slice().sort((a, b) => a.time.localeCompare(b.time));
  if (!list.length) return '<p class="empty">لا توجد أحداث مسجلة لهذا اليوم</p>';
  return list.map(e => `
    <div class="ev ${e.done ? 'done' : ''}">
      <input type="checkbox" data-done="${ds}|${esc(e.id)}" ${e.done ? 'checked' : ''} title="تم">
      <div class="ev-main">
        <span class="ev-title" dir="auto">${esc(e.title)}</span>
        <span class="ev-time">⏰ ${esc(fmtTime(e.time))}${e.remind ? ' · تنبيه قبلها بـ ' + e.remind + ' د' : ''}</span>
      </div>
      <button class="del" data-del="${ds}|${esc(e.id)}">حذف</button>
    </div>`).join('');
}

function dayCardHTML(d) {
  const ds = fmt(d), isToday = ds === fmt(T);
  return `<div class="day ${isToday ? 'today' : ''}">
    <div class="day-head">
      <span class="day-name">${DAYS[dayIdx(d)]}${isToday ? '<span class="badge">اليوم</span>' : ''}</span>
      <span class="date-pill">${short(d)}</span>
    </div>${eventsHTML(ds)}</div>`;
}

function renderMonth(days) {
  let h = '<div class="card"><div class="month-grid">';
  h += DAYS.map(n => `<div class="mh">${n}</div>`).join('');
  for (let i = 0; i < dayIdx(days[0]); i++) h += '<div class="cell blank"></div>';
  days.forEach(d => {
    const ds = fmt(d), list = (events[ds] || []).slice().sort((a, b) => a.time.localeCompare(b.time));
    h += `<button class="cell ${ds === fmt(T) ? 'today' : ''} ${ds === selected ? 'sel' : ''}" data-pick="${ds}">
      <span class="dn">${d.getDate()}</span>
      ${list.slice(0, 2).map(e => `<span class="chip ${e.done ? 'done' : ''}" dir="auto">${esc(e.time)} ${esc(e.title)}</span>`).join('')}
      ${list.length > 2 ? `<span class="more">+${list.length - 2} أخرى</span>` : ''}
    </button>`;
  });
  h += '</div></div><div class="days" style="margin-top:16px">' + dayCardHTML(parseDate(selected)) + '</div>';
  $('cal-body').innerHTML = h;
}

async function addEvent() {
  const title = $('event-input').value.trim();
  const time = $('time-input').value || '08:00';
  const date = $('day-select').value;
  const remind = parseInt($('remind-select').value, 10) || 0;
  if (!title) { toast('اكتب اسم الحدث أو المهمة أولاً', 'warn'); $('event-input').focus(); return; }
  if (!date) { toast('اختر اليوم', 'warn'); return; }
  (events[date] = events[date] || []).push({id: uid(), title, time, remind, done: false});
  selected = date;
  $('event-input').value = '';
  renderCal();
  await persistEvents();
  toast('تمت إضافة الحدث ✅', 'ok');
}

function findEvent(key) {
  const [ds, id] = key.split('|');
  const list = events[ds] || [];
  return {ds, list, idx: list.findIndex(e => e.id === id)};
}

$('cal-body').addEventListener('click', async ev => {
  const del = ev.target.closest('[data-del]');
  if (del) {
    const {ds, list, idx} = findEvent(del.dataset.del);
    if (idx > -1) {
      list.splice(idx, 1);
      if (!list.length) delete events[ds];
      renderCal();
      await persistEvents();
      toast('تم حذف الحدث', 'ok');
    }
    return;
  }
  const pick = ev.target.closest('[data-pick]');
  if (pick) { selected = pick.dataset.pick; $('day-select').value = selected; renderCal(); }
});
$('cal-body').addEventListener('change', async ev => {
  const c = ev.target.closest('[data-done]');
  if (!c) return;
  const {list, idx} = findEvent(c.dataset.done);
  if (idx > -1) { list[idx].done = c.checked; renderCal(); await persistEvents(); }
});

function shift(dir) {
  if (view === 'week') anchor = addDays(anchor, 7 * dir);
  else anchor = new Date(anchor.getFullYear(), anchor.getMonth() + dir, 1);
  renderCal();
}
$('prev-btn').onclick = () => shift(-1);
$('next-btn').onclick = () => shift(1);
$('today-btn').onclick = () => { anchor = new Date(T.getFullYear(), T.getMonth(), T.getDate()); selected = fmt(anchor); renderCal(); };
$('view-week').onclick = () => { view = 'week'; if (selected) anchor = parseDate(selected); renderCal(); };
$('view-month').onclick = () => { view = 'month'; const s = parseDate(selected); anchor = new Date(s.getFullYear(), s.getMonth(), 1); renderCal(); };
$('add-btn').onclick = addEvent;
$('event-input').addEventListener('keydown', e => { if (e.key === 'Enter') addEvent(); });
$('day-select').onchange = e => { selected = e.target.value; };

/* ---------- الإشعارات ---------- */
function refreshNotifBtn() {
  const b = $('notif-btn');
  if (!('Notification' in window)) { b.textContent = '🔕 غير مدعوم'; return; }
  if (Notification.permission === 'granted') { b.textContent = '🔔 الإشعارات مفعلة'; b.classList.add('on'); }
}
function requestNotif() {
  if (!('Notification' in window)) { toast('متصفحك لا يدعم الإشعارات', 'err'); return; }
  Notification.requestPermission().then(p => {
    if (p === 'granted') { toast('تم تفعيل الإشعارات ✅', 'ok'); refreshNotifBtn(); new Notification('Abdullah Master Workspace 🐉', {body: 'هتوصلك تنبيهات المهام وانت فاتح الصفحة'}); }
    else toast('تم رفض إذن الإشعارات', 'warn');
  });
}
const fired = new Set(JSON.parse(lsGet('fired') || '[]'));
function checkReminders() {
  const now = new Date(), ds = fmt(now);
  (events[ds] || []).forEach(e => {
    if (e.done) return;
    const [h, m] = e.time.split(':').map(Number);
    const at = new Date(now.getFullYear(), now.getMonth(), now.getDate(), h, m);
    const fireAt = at.getTime() - (e.remind || 0) * 60000;
    const key = ds + '|' + e.id + '|' + e.remind;
    if (!fired.has(key) && now.getTime() >= fireAt && now.getTime() <= at.getTime() + 15 * 60000) {
      fired.add(key);
      lsSet('fired', JSON.stringify([...fired].slice(-200)));
      const msg = e.remind ? `بعد ${e.remind} دقيقة: ${e.title}` : `حان موعد: ${e.title}`;
      if ('Notification' in window && Notification.permission === 'granted') new Notification('تنبيه من جدول عبدالله 🐉', {body: msg});
      toast('⏰ ' + msg, 'warn');
    }
  });
}
setInterval(checkReminders, 15000);

/* ================= السلايدات ================= */
const CHUNK = 4 * 1024 * 1024;
const EXTS = ['.pdf', '.docx', '.pptx', '.txt', '.md', '.csv'];
let upload = {id: null, token: 0, busy: false};

function fmtSize(b) { return b > 1048576 ? (b / 1048576).toFixed(1) + ' MB' : Math.max(1, Math.round(b / 1024)) + ' KB'; }
function setProgress(pct, text, kind) {
  $('progress-box').classList.remove('hidden');
  const bar = $('prog-bar');
  bar.style.width = Math.max(0, Math.min(100, pct)) + '%';
  bar.className = 'bar-fg' + (kind ? ' ' + kind : '');
  $('prog-pct').textContent = Math.round(pct) + '%';
  if (text) $('prog-text').textContent = text;
}
function setConvert(enabled, label) {
  const b = $('convert-btn');
  b.disabled = !enabled;
  b.textContent = label;
}

async function sendChunk(id, idx, blob) {
  for (let attempt = 0; attempt < 4; attempt++) {
    try {
      const r = await fetch(`/api/upload/chunk?id=${encodeURIComponent(id)}&index=${idx}`, {method: 'POST', body: blob, headers: {'Content-Type': 'application/octet-stream'}});
      if (r.ok) return;
      let j = {};
      try { j = await r.json(); } catch (e) {}
      if (r.status === 409 && j.expected === idx + 1) return;       // وصلت قبل كده
      const err = new Error(j.error || ('خطأ ' + r.status));
      if (r.status >= 400 && r.status < 500) err.fatal = true;
      throw err;
    } catch (e) {
      if (e.fatal || attempt === 3) throw e;
      await sleep(700 * (attempt + 1));
    }
  }
}

async function handleFile(input) {
  const f = input.files && input.files[0];
  const token = ++upload.token;
  if (upload.id) { postJSON('/api/upload/discard', {id: upload.id}).catch(() => {}); upload.id = null; }
  $('progress-box').classList.add('hidden');
  if (!f) { $('file-name').textContent = 'No file selected.'; setConvert(false, 'تحويل الملف إلى سلايدات 📤'); return; }
  $('file-name').textContent = f.name + ' (' + fmtSize(f.size) + ')';
  const ext = (f.name.match(/\.[^.]+$/) || [''])[0].toLowerCase();
  if (!EXTS.includes(ext)) { toast('نوع الملف غير مدعوم (PDF / DOCX / PPTX / TXT / MD / CSV)', 'err'); setConvert(false, 'تحويل الملف إلى سلايدات 📤'); return; }
  if (!f.size) { toast('الملف فارغ', 'err'); setConvert(false, 'تحويل الملف إلى سلايدات 📤'); return; }
  if (f.size > 2 * 1024 ** 3) { toast('الملف أكبر من 2GB', 'err'); return; }
  if (ext === '.pdf' && !CONFIG.pdf) toast('مكتبة PDF غير مثبتة على السيرفر: pip install pypdf', 'warn');
  if (ext === '.docx' && !CONFIG.docx) toast('مكتبة Word غير مثبتة: pip install python-docx', 'warn');
  if (ext === '.pptx' && !CONFIG.pptx) toast('مكتبة PowerPoint غير مثبتة: pip install python-pptx', 'warn');

  setConvert(false, 'جاري رفع الملف...');
  setProgress(0, 'جاري رفع الملف...');
  try {
    const st = await postJSON('/api/upload/start', {name: f.name, size: f.size});
    let sent = 0, idx = 0;
    while (sent < f.size) {
      if (token !== upload.token) { postJSON('/api/upload/discard', {id: st.id}).catch(() => {}); return; }
      const blob = f.slice(sent, sent + CHUNK);
      await sendChunk(st.id, idx, blob);
      sent += blob.size; idx++;
      setProgress(sent / f.size * 100, `جاري رفع الملف... ${fmtSize(sent)} من ${fmtSize(f.size)}`);
    }
    await postJSON('/api/upload/finish', {id: st.id});
    if (token !== upload.token) { postJSON('/api/upload/discard', {id: st.id}).catch(() => {}); return; }
    upload.id = st.id;
    setProgress(100, 'تم رفع الملف بنجاح ✅ — دوس تحويل', 'ok');
    setConvert(true, 'تحويل الملف إلى سلايدات 📤');
  } catch (e) {
    if (token !== upload.token) return;
    setProgress(parseFloat($('prog-pct').textContent) || 0, 'فشل الرفع: ' + e.message, 'err');
    setConvert(false, 'تحويل الملف إلى سلايدات 📤');
    toast('فشل الرفع: ' + e.message, 'err');
  }
}

async function convertFile() {
  if (!upload.id) { toast('ارفع ملف الأول', 'warn'); return; }
  const id = upload.id;
  setConvert(false, 'جاري التحويل...');
  setProgress(1, 'بدء المعالجة...');
  try {
    const {job} = await postJSON('/api/convert', {id, slides: parseInt($('slides-count').value, 10) || 0, ai: $('ai-check').checked, ocr: $('ocr-check').checked});
    while (true) {
      const r = await fetch('/api/job/' + job);
      const j = await r.json();
      if (!r.ok) throw new Error(j.error || 'فقدنا المهمة');
      if (j.status === 'error') throw new Error(j.error || 'فشلت المعالجة');
      setProgress(j.progress, j.message, j.status === 'done' ? 'ok' : '');
      if (j.status === 'done') {
        setConvert(true, 'تحويل الملف إلى سلايدات 📤');
        toast('تم توليد ' + j.result.slides.length + ' سلايد ✅', 'ok');
        if (j.warn) toast(j.warn, 'warn');
        openEditor(j.result.title, j.result.slides);
        return;
      }
      await sleep(700);
    }
  } catch (e) {
    setProgress(parseFloat($('prog-pct').textContent) || 0, 'فشل: ' + e.message, 'err');
    setConvert(true, 'تحويل الملف إلى سلايدات 📤');
    toast(e.message, 'err');
  }
}

/* ---- توليد من النص المكتوب ---- */
function localDeck(title, lines) {
  const slides = [{type: 'title', title, bullets: []}];
  const hasHeads = lines.some(l => l.startsWith('# '));
  if (hasHeads) {
    let cur = null;
    lines.forEach(l => {
      if (l.startsWith('# ')) { cur = {type: 'content', title: l.slice(2).trim(), bullets: []}; slides.push(cur); }
      else { if (!cur) { cur = {type: 'content', title, bullets: []}; slides.push(cur); } cur.bullets.push(l); }
    });
  } else {
    const groups = [];
    for (let i = 0; i < lines.length; i += 6) groups.push(lines.slice(i, i + 6));
    groups.forEach((g, i) => slides.push({type: 'content', title: groups.length > 1 ? `${title} — ${i + 1}` : title, bullets: g}));
  }
  return {title, slides};
}
async function generateFromText() {
  const title = $('slide-title').value.trim();
  const text = $('slide-points').value.trim();
  if (!title || !text) { toast('اكتب العنوان والنقاط أولاً', 'warn'); return; }
  const btn = $('gen-btn');
  btn.disabled = true;
  try {
    let deck;
    if ($('smart-check').checked && text.length > 1200) {
      deck = await postJSON('/api/slides/from_text', {title, text, slides: parseInt($('slides-count').value, 10) || 0});
    } else {
      deck = localDeck(title, text.split('\n').map(x => x.trim()).filter(Boolean));
    }
    openEditor(deck.title, deck.slides);
    toast('تم توليد ' + deck.slides.length + ' سلايد — صدّرها من الأعلى', 'ok');
  } catch (e) { toast(e.message, 'err'); }
  btn.disabled = false;
}

/* ================= شكل السلايد ================= */
function slideHTML(s, i, n) {
  const bullets = (s.bullets || []).filter(Boolean);
  if (s.type === 'title') {
    return `<div class="slide cover"><div class="em">🐉</div><h1 dir="auto">${esc(s.title)}</h1>${bullets.length ? `<p dir="auto">${esc(bullets.join(' · '))}</p>` : ''}<div class="num">${i + 1} / ${n}</div></div>`;
  }
  const len = bullets.join('').length + bullets.length * 12;
  const sz = len > 1000 ? 'xs' : len > 700 ? 'sm' : len > 430 ? 'md' : 'lg';
  return `<div class="slide ${sz}"><div class="accent"></div><h2 dir="auto">${esc(s.title)}</h2><ul>${bullets.map(b => `<li dir="auto">${esc(b)}</li>`).join('')}</ul><div class="num">${i + 1} / ${n}</div><div class="brand-s">Abdullah Master Workspace</div></div>`;
}
function fit(box) {
  const s = box.firstElementChild;
  if (s && box.clientWidth) s.style.transform = 'scale(' + (box.clientWidth / 1280) + ')';
}
const ro = new ResizeObserver(es => es.forEach(e => fit(e.target)));
function mountSlide(box, s, i, n) { box.innerHTML = slideHTML(s, i, n); fit(box); ro.observe(box); }

/* ================= المحرر ================= */
const ED = {title: '', slides: [], cur: 0};
let deckTimer = null;
function saveDeckSoon() {
  clearTimeout(deckTimer);
  deckTimer = setTimeout(() => {
    const deck = {title: ED.title, slides: ED.slides};
    postJSON('/api/data', {deck}).then(() => setLastDeck(deck)).catch(() => {});
  }, 900);
}
function setLastDeck(deck) {
  const b = $('last-deck-btn');
  b.classList.remove('hidden');
  b.textContent = `📂 فتح آخر عرض محفوظ: ${deck.title} (${deck.slides.length} سلايد)`;
  b.onclick = () => openEditor(deck.title, deck.slides);
}
const normSlide = s => ({type: s.type === 'title' ? 'title' : 'content', title: s.title || '', bullets: (s.bullets || []).map(String)});

function openEditor(title, slides) {
  ED.title = title || 'عرض';
  ED.slides = slides.map(normSlide);
  if (!ED.slides.length) ED.slides = [{type: 'title', title: ED.title, bullets: []}];
  ED.cur = 0;
  $('editor').classList.remove('hidden');
  document.body.classList.add('no-scroll');
  $('ed-title').value = ED.title;
  renderList();
  loadForm();
  saveDeckSoon();
}
function closeEditor() { $('editor').classList.add('hidden'); document.body.classList.remove('no-scroll'); }

function renderList() {
  $('ed-list').innerHTML = ED.slides.map((s, i) => `<button class="thumb ${i === ED.cur ? 'active' : ''}" data-i="${i}"><span class="n">${i + 1}</span><span class="tt" dir="auto">${esc(s.title || '(بدون عنوان)')}</span></button>`).join('');
  const a = $('ed-list').querySelector('.active');
  if (a) a.scrollIntoView({block: 'nearest'});
}
function loadForm() {
  const s = ED.slides[ED.cur];
  $('ed-slide-title').value = s.title;
  $('ed-slide-bullets').value = s.bullets.join('\n');
  mountSlide($('ed-preview'), s, ED.cur, ED.slides.length);
}
function refreshPreview() { mountSlide($('ed-preview'), ED.slides[ED.cur], ED.cur, ED.slides.length); }

$('ed-list').addEventListener('click', e => {
  const b = e.target.closest('[data-i]');
  if (b) { ED.cur = +b.dataset.i; renderList(); loadForm(); }
});
$('ed-title').addEventListener('input', e => {
  ED.title = e.target.value;
  if (ED.slides[0] && ED.slides[0].type === 'title') { ED.slides[0].title = ED.title; if (ED.cur === 0) { $('ed-slide-title').value = ED.title; refreshPreview(); } renderList(); }
  saveDeckSoon();
});
$('ed-slide-title').addEventListener('input', e => {
  ED.slides[ED.cur].title = e.target.value;
  if (ED.cur === 0 && ED.slides[0].type === 'title') { ED.title = e.target.value; $('ed-title').value = e.target.value; }
  refreshPreview(); renderList(); saveDeckSoon();
});
$('ed-slide-bullets').addEventListener('input', e => {
  ED.slides[ED.cur].bullets = e.target.value.split('\n').map(x => x.trim()).filter(Boolean);
  refreshPreview(); saveDeckSoon();
});
$('btn-add-slide').onclick = () => {
  ED.slides.splice(ED.cur + 1, 0, {type: 'content', title: 'سلايد جديد', bullets: ['']});
  ED.slides[ED.cur + 1].bullets = [];
  ED.cur++; renderList(); loadForm(); saveDeckSoon();
  $('ed-slide-title').focus(); $('ed-slide-title').select();
};
$('btn-dup').onclick = () => {
  const c = JSON.parse(JSON.stringify(ED.slides[ED.cur]));
  c.type = 'content';
  ED.slides.splice(ED.cur + 1, 0, c); ED.cur++; renderList(); loadForm(); saveDeckSoon();
};
$('btn-del').onclick = () => {
  if (ED.slides.length <= 1) { toast('لازم يفضل سلايد واحد على الأقل', 'warn'); return; }
  ED.slides.splice(ED.cur, 1);
  ED.cur = Math.min(ED.cur, ED.slides.length - 1);
  renderList(); loadForm(); saveDeckSoon();
};
function moveSlide(d) {
  const j = ED.cur + d;
  if (j < 0 || j >= ED.slides.length) return;
  [ED.slides[ED.cur], ED.slides[j]] = [ED.slides[j], ED.slides[ED.cur]];
  ED.cur = j; renderList(); loadForm(); saveDeckSoon();
}
$('btn-up').onclick = () => moveSlide(-1);
$('btn-down').onclick = () => moveSlide(1);
$('btn-ed-close').onclick = closeEditor;

/* ---- التصدير ---- */
function downloadBlob(blob, name) {
  const a = document.createElement('a');
  a.href = URL.createObjectURL(blob); a.download = name;
  document.body.appendChild(a); a.click();
  setTimeout(() => { URL.revokeObjectURL(a.href); a.remove(); }, 1500);
}
const safeName = s => (s || 'slides').replace(/[\\/:*?"<>|]+/g, ' ').trim().slice(0, 80) || 'slides';

async function exportPPTX() {
  const b = $('btn-pptx'); b.disabled = true;
  try {
    const r = await fetch('/api/export/pptx', {method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({title: ED.title, slides: ED.slides})});
    if (!r.ok) { let j = {}; try { j = await r.json(); } catch (e) {} throw new Error(j.error || ('خطأ ' + r.status)); }
    downloadBlob(await r.blob(), safeName(ED.title) + '.pptx');
    toast('تم تنزيل ملف PowerPoint ✅', 'ok');
  } catch (e) { toast(e.message, 'err'); }
  b.disabled = false;
}
async function exportPDF() {
  const root = $('print-root');
  root.innerHTML = ED.slides.map((s, i) => `<div class="print-page">${slideHTML(s, i, ED.slides.length)}</div>`).join('');
  const oldTitle = document.title;
  document.title = safeName(ED.title);
  try { await document.fonts.ready; } catch (e) {}
  toast('في نافذة الطباعة اختر "حفظ كـ PDF" (Save as PDF)', 'ok');
  await sleep(400);
  window.print();
  setTimeout(() => { document.title = oldTitle; }, 1000);
}
window.addEventListener('afterprint', () => { $('print-root').innerHTML = ''; });
$('btn-pptx').onclick = exportPPTX;
$('btn-pdf').onclick = exportPDF;

/* ---- العرض التقديمي ---- */
let pres = {i: 0};
function sizePresent() {
  const box = $('present-box');
  const w = Math.min(window.innerWidth, window.innerHeight * 16 / 9);
  box.style.width = w + 'px';
  fit(box);
}
function showPres() {
  const n = ED.slides.length;
  pres.i = Math.max(0, Math.min(n - 1, pres.i));
  mountSlide($('present-box'), ED.slides[pres.i], pres.i, n);
  $('pres-count').textContent = (pres.i + 1) + ' / ' + n;
  sizePresent();
}
function startPresent() {
  pres.i = ED.cur;
  $('present').classList.remove('hidden');
  showPres();
  const el = $('present');
  if (el.requestFullscreen) el.requestFullscreen().catch(() => {});
}
function endPresent() {
  $('present').classList.add('hidden');
  if (document.fullscreenElement && document.exitFullscreen) document.exitFullscreen().catch(() => {});
}
$('btn-present').onclick = startPresent;
$('pres-next').onclick = e => { e.stopPropagation(); pres.i++; showPres(); };
$('pres-prev').onclick = e => { e.stopPropagation(); pres.i--; showPres(); };
$('pres-exit').onclick = e => { e.stopPropagation(); endPresent(); };
$('present').addEventListener('click', e => { if (!e.target.closest('.pres-ui')) { pres.i++; showPres(); } });
window.addEventListener('resize', () => { if (!$('present').classList.contains('hidden')) sizePresent(); });
document.addEventListener('fullscreenchange', () => { if (!document.fullscreenElement) $('present').classList.add('hidden'); else sizePresent(); });
document.addEventListener('keydown', e => {
  if ($('present').classList.contains('hidden')) return;
  if (['ArrowLeft', 'PageDown', ' ', 'Enter', 'ArrowDown'].includes(e.key)) { e.preventDefault(); pres.i++; showPres(); }
  else if (['ArrowRight', 'PageUp', 'Backspace', 'ArrowUp'].includes(e.key)) { e.preventDefault(); pres.i--; showPres(); }
  else if (e.key === 'Escape') endPresent();
});

/* ================= متابعة المنهج ================= */
const STUDY_SUBJECTS = {
            salah: {
                id: 'salah',
                title: 'عربي: م. محمد صلاح',
                icon: '📖',
                groups: [
                    {
                        name: '📚 النصوص',
                        lessons: [
                            { id: 'salah_txt_1', title: 'الوحدة 1 - معلوماتي: العلم بين القانون والضمير' },
                            { id: 'salah_txt_2', title: 'الوحدة 1 - نثري: جوهر الفضيلة' },
                            { id: 'salah_txt_3', title: 'الوحدة 1 - شعري: العلم نور' },
                            { id: 'salah_txt_4', title: 'الوحدة 2 - معلوماتي: التقنية وصناعة الرأي العام' },
                            { id: 'salah_txt_5', title: 'الوحدة 2 - نثري: بين الذوق والعقل' },
                            { id: 'salah_txt_6', title: 'الوحدة 2 - شعري: أثر الكلمة وحفظ اللسان' },
                            { id: 'salah_txt_7', title: 'الوحدة 3 - معلوماتي: حقوق الإنسان بين العالمية والخصوصية' },
                            { id: 'salah_txt_8', title: 'الوحدة 3 - نثري: الحرية الفكرية' },
                            { id: 'salah_txt_9', title: 'الوحدة 3 - شعري: يا أخي الإنسان' }
                        ]
                    },
                    {
                        name: '📝 النحو (ما سبق دراسته)',
                        lessons: [
                            { id: 'salah_gram_p1', title: 'علامات الإعراب' },
                            { id: 'salah_gram_p2', title: 'الجملة الفعلية' },
                            { id: 'salah_gram_p3', title: 'اللازم والمتعدي - المفعول به' },
                            { id: 'salah_gram_p4', title: 'بناء الفعل للمجهول ونائب الفاعل' },
                            { id: 'salah_gram_p5', title: 'الحال - المفعول لأجله' },
                            { id: 'salah_gram_p6', title: 'الظرف' },
                            { id: 'salah_gram_p7', title: 'المفعول المطلق' },
                            { id: 'salah_gram_p8', title: 'الجملة الاسمية (المبتدأ والخبر)' },
                            { id: 'salah_gram_p9', title: 'كان وأخواتها' },
                            { id: 'salah_gram_p10', title: 'كاد وأخواتها' },
                            { id: 'salah_gram_p11', title: 'إن وأخواتها' },
                            { id: 'salah_gram_p12', title: 'أسلوب المدح والذم' },
                            { id: 'salah_gram_p13', title: 'الملحقات' },
                            { id: 'salah_gram_p14', title: 'كم الاستفهامية وكم الخبرية' }
                        ]
                    },
                    {
                        name: '🎯 النحو المقرر',
                        lessons: [
                            { id: 'salah_gram_r1', title: 'الأسماء الخمسة' },
                            { id: 'salah_gram_r2', title: 'الممنوع من الصرف' },
                            { id: 'salah_gram_r3', title: 'النعت' },
                            { id: 'salah_gram_r4', title: 'البدل' },
                            { id: 'salah_gram_r5', title: 'العطف' },
                            { id: 'salah_gram_r6', title: 'همزة القطع وألف الوصل' },
                            { id: 'salah_gram_r7', title: 'الهمزة المتوسطة' },
                            { id: 'salah_gram_r8', title: 'الهمزة المتطرفة' },
                            { id: 'salah_gram_r9', title: 'الفرق بين التاء المربوطة والهاء والمفتوحة' }
                        ]
                    },
                    {
                        name: '✨ البلاغة',
                        lessons: [
                            { id: 'salah_bal_1', title: 'التشبيه' },
                            { id: 'salah_bal_2', title: 'الاستعارة' },
                            { id: 'salah_bal_3', title: 'أنواع الصور الخيالية' },
                            { id: 'salah_bal_4', title: 'المحسنات البديعية' },
                            { id: 'salah_bal_5', title: 'الأسلوب الخبري' },
                            { id: 'salah_bal_6', title: 'الأسلوب الإنشائي وأنواعه' },
                            { id: 'salah_bal_7', title: 'الإيجاز' },
                            { id: 'salah_bal_8', title: 'الإطناب' },
                            { id: 'salah_bal_9', title: 'الموسيقى' },
                            { id: 'salah_bal_10', title: 'الوحدة الفنية' },
                            { id: 'salah_bal_11', title: 'العلاقات بين الجمل' }
                        ]
                    },
                    {
                        name: '✍️ التعبير',
                        lessons: [
                            { id: 'salah_exp_1', title: 'الفقرة الحجاجية' },
                            { id: 'salah_exp_2', title: 'حوار قصير' },
                            { id: 'salah_exp_3', title: 'موقف نقدي' },
                            { id: 'salah_exp_4', title: 'كتابة فقرة نقدية' },
                            { id: 'salah_exp_5', title: 'كتابة رأي' },
                            { id: 'salah_exp_6', title: 'الكتابة الموجهة' },
                            { id: 'salah_exp_7', title: 'مشهد سري موجز' },
                            { id: 'salah_exp_8', title: 'إعادة صياغة فقرة بإيجاز' },
                            { id: 'salah_exp_9', title: 'كتابة موقف حياتي' },
                            { id: 'salah_exp_10', title: 'كتابة لوحة إرشادية' },
                            { id: 'salah_exp_11', title: 'كتابة فقرة استهلالية' },
                            { id: 'salah_exp_12', title: 'كتابة فقرة تحليلية' },
                            { id: 'salah_exp_13', title: 'كتابة مقطع قصصي قصير' }
                        ]
                    }
                ]
            },
            hajjaj: {
                id: 'hajjaj',
                title: 'عربي: م. أبو الحجاج',
                icon: '🪶',
                groups: [
                    {
                        name: 'الوحدة 1: العلم والضمير',
                        lessons: [
                            { id: 'hajjaj_1_1', title: 'الدرس الأول: العلم والضمير (قواعد: الأسماء الخمسة)' },
                            { id: 'hajjaj_1_2', title: 'الدرس الثاني: الخلق جوهر الفضيلة (قواعد: تابع الأسماء الخمسة)' },
                            { id: 'hajjaj_1_3', title: 'الدرس الثالث: (نص) العلم نور (تطبيق على الأسماء الخمسة)' }
                        ]
                    },
                    {
                        name: 'الوحدة 2: سلطان الكلمة الرقمية',
                        lessons: [
                            { id: 'hajjaj_2_1', title: 'الدرس الأول: التقنية وصناعة الرأي العام (قواعد: الممنوع من الصرف)' },
                            { id: 'hajjaj_2_2', title: 'الدرس الثاني: وصية إلى ولدي (قواعد: تابع الممنوع من الصرف)' },
                            { id: 'hajjaj_2_3', title: 'الدرس الثالث: (نص) أثر الكلمة وحفظ اللسان' }
                        ]
                    },
                    {
                        name: 'الوحدة 3: الإنسان وكرامته',
                        lessons: [
                            { id: 'hajjaj_3_1', title: 'الدرس الأول: حقوق الإنسان بين العالمية والخصوصية (قواعد: النعت والبدل والتمييز)' },
                            { id: 'hajjaj_3_2', title: 'الدرس الثاني: الحرية الفكرية (قواعد: العطف)' },
                            { id: 'hajjaj_3_3', title: 'الدرس الثالث: (نص) يا أخي الإنسان' }
                        ]
                    }
                ]
            },
            history: {
                id: 'history',
                title: 'مادة التاريخ',
                icon: '🏛️',
                groups: [
                    {
                        name: 'الوحدة 1: مصر بين الماضي والحاضر',
                        lessons: [
                            { id: 'his_1_1', title: 'الدرس الأول: بناء الدولة المصرية واستمرارها عبر التاريخ' },
                            { id: 'his_1_2', title: 'الدرس الثاني: الدولة المصرية القديمة ومقاومة الاحتلال' },
                            { id: 'his_1_3', title: 'الدرس الثالث: التحولات الكبرى في مصر (العصر الوسيط)' },
                            { id: 'his_1_4', title: 'الدرس الرابع: التحولات الكبرى في مصر (العصر الحديث)' }
                        ]
                    },
                    {
                        name: 'الوحدة 2: تحولات القوة وبناء الوعي بمصر',
                        lessons: [
                            { id: 'his_2_1', title: 'الدرس الأول: سياسات الاحتلال بمصر في العصر الحديث' },
                            { id: 'his_2_2', title: 'الدرس الثاني: صمود الشعب المصري ومقاومة الاحتلال' },
                            { id: 'his_2_3', title: 'الدرس الثالث: ثورة 23 يوليو 1952م والتحولات الكبرى' },
                            { id: 'his_2_4', title: 'الدرس الرابع: الهوية الوطنية في مواجهة الاستعمار الجديد' }
                        ]
                    },
                    {
                        name: 'الوحدة 3: مصر وقضايا التحرر الوطني',
                        lessons: [
                            { id: 'his_3_1', title: 'الدرس الأول: حركات التحرر الوطني ضد الاحتلال الفرنسي والإيطالي' },
                            { id: 'his_3_2', title: 'الدرس الثاني: حركات التحرر الوطني ضد الاحتلال البريطاني والإسباني' },
                            { id: 'his_3_3', title: 'الدرس الثالث: مرتكزات الدور المصري إقليمياً ودولياً' },
                            { id: 'his_3_4', title: 'الدرس الرابع: دور مصر في دعم حركات التحرر الوطني' }
                        ]
                    }
                ]
            },
            english: {
                id: 'english',
                title: 'مادة الإنجليزية (English)',
                icon: '🔤',
                groups: [
                    {
                        name: 'Unit 1',
                        lessons: [
                            { id: 'eng_u1_vocab', title: 'Vocabulary & Reading' },
                            { id: 'eng_u1_gram', title: 'Grammar' }
                        ]
                    },
                    {
                        name: 'Unit 2',
                        lessons: [
                            { id: 'eng_u2_vocab', title: 'Vocabulary & Reading' },
                            { id: 'eng_u2_gram', title: 'Grammar' }
                        ]
                    },
                    {
                        name: 'Unit 3',
                        lessons: [
                            { id: 'eng_u3_vocab', title: 'Vocabulary & Reading' },
                            { id: 'eng_u3_gram', title: 'Grammar' }
                        ]
                    },
                    {
                        name: 'Unit 4',
                        lessons: [
                            { id: 'eng_u4_vocab', title: 'Vocabulary & Reading' },
                            { id: 'eng_u4_gram', title: 'Grammar' }
                        ]
                    },
                    {
                        name: 'Unit 5',
                        lessons: [
                            { id: 'eng_u5_vocab', title: 'Vocabulary & Reading' },
                            { id: 'eng_u5_gram', title: 'Grammar' }
                        ]
                    },
                    {
                        name: 'Unit 6',
                        lessons: [
                            { id: 'eng_u6_vocab', title: 'Vocabulary & Reading' },
                            { id: 'eng_u6_gram', title: 'Grammar' }
                        ]
                    }
                ]
            }
        };
const STUDY = {progress: {}, custom: {}, tab: 'all'};
let studySaving = Promise.resolve();
function persistStudy() {
  const payload = {study: {progress: STUDY.progress, custom: STUDY.custom}};
  studySaving = studySaving.then(() => postJSON('/api/data', payload)).catch(() => toast('تعذر حفظ التقدم', 'err'));
}
function subjLessons(key) {
  const s = STUDY_SUBJECTS[key];
  if (!s) return [];
  return s.groups.flatMap(g => g.lessons).concat(STUDY.custom[key] || []);
}
function subjStats(key) {
  const l = subjLessons(key), done = l.filter(x => STUDY.progress[x.id]).length;
  return {total: l.length, done, pct: l.length ? Math.round(done / l.length * 100) : 0};
}
function lessonHTML(l, key, custom) {
  const on = !!STUDY.progress[l.id];
  return `<div class="lesson ${on ? 'done' : ''}"><label style="display:flex;align-items:center;gap:12px;flex:1;cursor:pointer;min-width:0">
      <input type="checkbox" data-lesson="${esc(l.id)}" ${on ? 'checked' : ''}><span class="lt" dir="auto">${esc(l.title)}</span></label>
      ${custom ? `<button class="del" data-st-del="${esc(key)}|${esc(l.id)}" title="حذف">حذف</button>` : ''}</div>`;
}
function renderStudy() {
  const keys = Object.keys(STUDY_SUBJECTS);
  if (STUDY.tab !== 'all' && !STUDY_SUBJECTS[STUDY.tab]) STUDY.tab = 'all';
  const typed = {};
  document.querySelectorAll('[data-st-input]').forEach(i => { typed[i.dataset.stInput] = i.value; });
  const y = window.scrollY;

  let total = 0, done = 0;
  keys.forEach(k => { const s = subjStats(k); total += s.total; done += s.done; });
  const pct = total ? Math.round(done / total * 100) : 0;
  $('st-pct').textContent = pct + '%';
  $('st-count').textContent = `${done} من ${total} درس`;
  $('st-bar').style.width = pct + '%';

  $('st-tabs').innerHTML = `<button class="chip-btn ${STUDY.tab === 'all' ? 'active' : ''}" data-st-tab="all">🗂️ الكل (الرئيسية)</button>` +
    keys.map(k => `<button class="chip-btn ${STUDY.tab === k ? 'active' : ''}" data-st-tab="${k}">${STUDY_SUBJECTS[k].icon} ${esc(STUDY_SUBJECTS[k].title)} <span class="st-badge">${subjStats(k).pct}%</span></button>`).join('');

  const show = STUDY.tab === 'all' ? keys : [STUDY.tab];
  $('st-views').innerHTML = show.map(k => {
    const s = STUDY_SUBJECTS[k], st = subjStats(k), cust = STUDY.custom[k] || [];
    return `<div class="card">
      <div class="subj-head">
        <div class="subj-title"><div class="subj-ico">${s.icon}</div><div><h2>${esc(s.title)}</h2><p>إجمالي الدروس: ${st.total} درس</p></div></div>
        <div class="subj-pct"><small>نسبة الإنجاز</small><b dir="ltr">${st.done} / ${st.total} (${st.pct}%)</b></div>
      </div>
      ${s.groups.map(g => `<div class="grp"><h4>${esc(g.name)}</h4>${g.lessons.map(l => lessonHTML(l, k, false)).join('')}</div>`).join('')}
      ${cust.length ? `<div class="grp custom"><h4>📌 دروس وملاحظات مضافة</h4>${cust.map(l => lessonHTML(l, k, true)).join('')}</div>` : ''}
      <div class="add-lesson"><h5>➕ إضافة درس أو ملاحظة خاصة لـ (${esc(s.title)}):</h5>
        <div class="r"><input type="text" data-st-input="${k}" maxlength="300" placeholder="عنوان الدرس الجديد...">
        <button class="btn sky" data-st-add="${k}">إضافة</button></div></div>
    </div>`;
  }).join('');

  document.querySelectorAll('[data-st-input]').forEach(i => { if (typed[i.dataset.stInput]) i.value = typed[i.dataset.stInput]; });
  window.scrollTo(0, y);
}
function addCustomLesson(key) {
  const input = document.querySelector(`[data-st-input="${key}"]`);
  const title = input ? input.value.trim() : '';
  if (!title) { toast('اكتب عنوان الدرس أولاً', 'warn'); if (input) input.focus(); return; }
  (STUDY.custom[key] = STUDY.custom[key] || []).push({id: `custom_${key}_${uid()}`, title});
  input.value = '';
  persistStudy();
  renderStudy();
  const again = document.querySelector(`[data-st-input="${key}"]`);
  if (again) again.focus();
}
$('tab-study').addEventListener('change', e => {
  const c = e.target.closest('[data-lesson]');
  if (!c) return;
  if (c.checked) STUDY.progress[c.dataset.lesson] = true; else delete STUDY.progress[c.dataset.lesson];
  persistStudy();
  renderStudy();
});
$('tab-study').addEventListener('click', e => {
  const t = e.target.closest('[data-st-tab]');
  if (t) { STUDY.tab = t.dataset.stTab; renderStudy(); return; }
  const a = e.target.closest('[data-st-add]');
  if (a) { addCustomLesson(a.dataset.stAdd); return; }
  const d = e.target.closest('[data-st-del]');
  if (d) {
    const [key, id] = d.dataset.stDel.split('|');
    STUDY.custom[key] = (STUDY.custom[key] || []).filter(l => l.id !== id);
    if (!STUDY.custom[key].length) delete STUDY.custom[key];
    delete STUDY.progress[id];
    persistStudy();
    renderStudy();
  }
});
$('tab-study').addEventListener('keydown', e => {
  if (e.key === 'Enter' && e.target.matches('[data-st-input]')) addCustomLesson(e.target.dataset.stInput);
});

/* ---------- ربط أزرار السلايدات + البداية ---------- */
$('upload-file').addEventListener('change', e => handleFile(e.target));
$('convert-btn').onclick = convertFile;
$('gen-btn').onclick = generateFromText;

(async function init() {
  try { CONFIG = await (await fetch('/api/config')).json(); } catch (e) {}
  $('ai-wrap').style.visibility = CONFIG.ai ? 'visible' : 'hidden';
  $('ocr-row').classList.toggle('hidden', !CONFIG.ocr);
  $('ocr-hint').classList.toggle('hidden', !!CONFIG.ocr);
  refreshNotifBtn();
  switchTab(lsGet('tab'));
  await loadData();
})();
</script>
</body>
</html>
"""


if __name__ == "__main__":
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "5000"))
    print(f"\n🐉 Abdullah Master Workspace على  http://localhost:{port}\n")
    app.run(host=host, port=port, debug=False, threaded=True)
